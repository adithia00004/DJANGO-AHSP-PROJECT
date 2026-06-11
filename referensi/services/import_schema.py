"""
Canonical Excel interchange contract for AHSP import/export.

This service owns the file-level schema used when AHSP rows cross module
boundaries through Excel. Legacy Excel layouts are still handled by adapters in
the view layer, but new exports should use this contract.
"""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any, BinaryIO, Iterable

from openpyxl import Workbook, load_workbook

from referensi.services.ahsp_code import normalize_ahsp_code
from referensi.services.import_repair import normalize_segment


SCHEMA_VERSION = "ahsp-interchange-1.0"
DATA_SHEET = "Data"
META_SHEET = "Meta"

REQUIRED_COLUMNS = (
    "kode_ahsp",
    "nama_ahsp",
    "segmen",
    "kode_item",
    "uraian",
    "satuan",
    "koefisien",
)

OPTIONAL_COLUMNS = (
    "no",
    "status",
    "alasan",
)

ALL_COLUMNS = REQUIRED_COLUMNS + OPTIONAL_COLUMNS


class UnsupportedImportSchema(ValueError):
    """Raised when a workbook is not AHSP Interchange v1."""


@dataclass(frozen=True)
class ImportSchemaRow:
    kode_ahsp: str
    nama_ahsp: str
    segmen: str
    kode_item: str
    uraian: str
    satuan: str
    koefisien: str
    no: str = ""
    status: str = ""
    alasan: str = ""

    def as_dict(self) -> dict[str, str]:
        return {
            "kode_ahsp": self.kode_ahsp,
            "nama_ahsp": self.nama_ahsp,
            "segmen": self.segmen,
            "kode_item": self.kode_item,
            "uraian": self.uraian,
            "satuan": self.satuan,
            "koefisien": self.koefisien,
            "no": self.no,
            "status": self.status,
            "alasan": self.alasan,
        }


def _clean(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower() in {"nan", "none"}:
        return ""
    return text


def _normalize_koef(value: Any) -> str:
    text = _clean(value).replace(",", ".")
    if not text:
        return "0"
    try:
        return str(Decimal(text))
    except (InvalidOperation, ValueError):
        return "0"


def _normalise_schema_row(row: dict[str, Any]) -> ImportSchemaRow:
    kode_ahsp = normalize_ahsp_code(_clean(row.get("kode_ahsp")))
    segmen = normalize_segment(row.get("segmen"))
    return ImportSchemaRow(
        kode_ahsp=kode_ahsp,
        nama_ahsp=_clean(row.get("nama_ahsp")) or kode_ahsp,
        segmen=segmen,
        kode_item=_clean(row.get("kode_item")) or "-",
        uraian=_clean(row.get("uraian")),
        satuan=_clean(row.get("satuan")) or "-",
        koefisien=_normalize_koef(row.get("koefisien")),
        no=_clean(row.get("no")),
        status=_clean(row.get("status")),
        alasan=_clean(row.get("alasan")),
    )


def dump_workbook(rows: Iterable[dict[str, Any]], meta: dict[str, Any] | None = None) -> bytes:
    """
    Serialize rows to AHSP Interchange v1.

    `sumber` is batch-level metadata and is stored in Meta, not repeated per row.
    """
    meta = meta or {}
    wb = Workbook()
    ws = wb.active
    ws.title = DATA_SHEET
    ws.append(list(ALL_COLUMNS))

    for raw in rows:
        row = _normalise_schema_row(raw)
        values = row.as_dict()
        ws.append([values.get(col, "") for col in ALL_COLUMNS])

    meta_ws = wb.create_sheet(META_SHEET)
    meta_ws.append(["key", "value"])
    meta_ws.append(["schema_version", SCHEMA_VERSION])
    meta_ws.append(["sumber", _clean(meta.get("sumber"))])
    meta_ws.append(["export_type", _clean(meta.get("export_type"))])

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def _read_meta(wb) -> dict[str, str]:
    if META_SHEET not in wb.sheetnames:
        return {}
    meta = {}
    ws = wb[META_SHEET]
    for key, value in ws.iter_rows(min_row=2, values_only=True):
        key_text = _clean(key)
        if key_text:
            meta[key_text] = _clean(value)
    return meta


def is_interchange_workbook(file_obj: BinaryIO) -> bool:
    pos = file_obj.tell() if hasattr(file_obj, "tell") else None
    try:
        wb = load_workbook(file_obj, read_only=True, data_only=True)
        meta = _read_meta(wb)
        if meta.get("schema_version") == SCHEMA_VERSION:
            return True
        if DATA_SHEET in wb.sheetnames:
            ws = wb[DATA_SHEET]
            header = [
                _clean(cell).lower()
                for cell in next(ws.iter_rows(min_row=1, max_row=1, values_only=True), [])
            ]
            return all(col in header for col in REQUIRED_COLUMNS)
        return False
    except Exception:
        return False
    finally:
        if pos is not None:
            file_obj.seek(pos)


def load_workbook_rows(file_obj: BinaryIO) -> tuple[list[dict[str, str]], dict[str, str]]:
    """
    Load AHSP Interchange v1 rows by column name.

    Returns `(rows, meta)`. Raises `UnsupportedImportSchema` when the workbook
    does not advertise/fit this schema.
    """
    wb = load_workbook(file_obj, read_only=True, data_only=True)
    meta = _read_meta(wb)
    if DATA_SHEET not in wb.sheetnames:
        raise UnsupportedImportSchema("Sheet Data tidak ditemukan.")

    ws = wb[DATA_SHEET]
    header = [
        _clean(cell).lower()
        for cell in next(ws.iter_rows(min_row=1, max_row=1, values_only=True), [])
    ]
    missing = [col for col in REQUIRED_COLUMNS if col not in header]
    if missing:
        raise UnsupportedImportSchema("Kolom wajib tidak lengkap: " + ", ".join(missing))

    if meta.get("schema_version") and meta.get("schema_version") != SCHEMA_VERSION:
        raise UnsupportedImportSchema(f"schema_version tidak didukung: {meta.get('schema_version')}")

    col_index = {name: idx for idx, name in enumerate(header)}
    rows: list[dict[str, str]] = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        raw = {col: values[col_index[col]] if col_index[col] < len(values) else "" for col in ALL_COLUMNS if col in col_index}
        normalized = _normalise_schema_row(raw)
        if not normalized.kode_ahsp or not normalized.uraian:
            continue
        if normalized.segmen not in {"A", "B", "C", "LAIN"}:
            continue
        rows.append(normalized.as_dict())

    meta.setdefault("schema_version", SCHEMA_VERSION)
    return rows, meta
