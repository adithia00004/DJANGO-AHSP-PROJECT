"""Paket perencanaan — 4 dokumen dalam satu berkas.

Isi paket WAJIB memakai adapter yang sama dengan unduhan tunggal; kalau tidak,
paket dan dokumen satuan bisa menyimpang tanpa ada yang menyadarinya.
"""

from datetime import date
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse

from dashboard.models import Project


class PaketPerencanaanTests(TestCase):
    def setUp(self):
        self.owner = get_user_model().objects.create_user("paket-owner", password="x")
        self.other = get_user_model().objects.create_user("paket-other", password="x")
        self.project = Project.objects.create(
            owner=self.owner,
            nama="Proyek Paket",
            sumber_dana="APBD",
            lokasi_project="Mataram",
            nama_client="Dinas Peternakan",
            nama_konsultan_perencana="Rozan Fahriady",
            anggaran_owner=Decimal("1000000"),
            tanggal_mulai=date(2026, 1, 1),
        )

    def _url(self, fmt):
        return reverse(
            "detail_project:export_paket_perencanaan",
            kwargs={"project_id": self.project.id, "format_type": fmt},
        )

    def test_package_contains_all_four_planning_documents(self):
        from detail_project.exports.export_manager import ExportManager

        documents = ExportManager(self.project)._collect_paket_documents()

        self.assertEqual(
            [d["title"] for d in documents],
            [title for title, _ in ExportManager.PAKET_PERENCANAAN],
        )
        for entry in documents:
            with self.subTest(title=entry["title"]):
                self.assertTrue(entry["data"], f"data {entry['title']} kosong")

    def test_every_format_returns_a_single_file(self):
        self.client.force_login(self.owner)
        for fmt in ("pdf", "word", "xlsx"):
            with self.subTest(fmt=fmt):
                response = self.client.get(self._url(fmt))
                self.assertEqual(response.status_code, 200)
                self.assertTrue(response["Content-Disposition"].startswith("attachment"))
                self.assertGreater(len(response.content), 5000)

    def test_unknown_format_is_rejected(self):
        self.client.force_login(self.owner)
        response = self.client.get(self._url("csv"))

        self.assertEqual(response.status_code, 400)

    def test_non_owner_cannot_download_another_projects_package(self):
        """Paket memuat seluruh nilai proyek; 404 menyembunyikan keberadaannya."""
        self.client.force_login(self.other)
        response = self.client.get(self._url("pdf"))

        self.assertEqual(response.status_code, 404)

    def test_package_reuses_single_download_data(self):
        """Paket tidak boleh punya jalur data sendiri yang bisa menyimpang."""
        from detail_project.exports.export_manager import ExportManager

        manager = ExportManager(self.project)
        paket_rekap = {
            d["title"]: d["data"] for d in manager._collect_paket_documents()
        }["REKAPITULASI RENCANA ANGGARAN BIAYA"]

        self.assertEqual(paket_rekap, manager._build_rekap_rab_data())

    def test_excel_package_has_no_invalid_formulas(self):
        """Excel menolak berkas berisi formula tak sah dan MEMBUANG isinya.

        Kolom Formula pada Volume berisi teks audit seperti "=Luas + Panjang".
        openpyxl memperlakukan string berawalan "=" sebagai formula, sehingga
        bila paket memakai export() generik alih-alih jalur XLSX khusus Volume,
        Excel menampilkan "Removed Records: Formula" dan membuang isinya.
        """
        import re
        import zipfile
        from io import BytesIO

        self.client.force_login(self.owner)
        response = self.client.get(self._url("xlsx"))
        self.assertEqual(response.status_code, 200)

        archive = zipfile.ZipFile(BytesIO(response.content))
        valid_start = re.compile(r"^('?[A-Za-z0-9 _]+'?!|SUM|IF|ROUND|ABS|[A-Z]+[0-9])")
        invalid = []
        for name in archive.namelist():
            if not name.startswith("xl/worksheets/sheet"):
                continue
            xml = archive.read(name).decode("utf-8", "ignore")
            for formula in re.findall(r"<f[^>]*>([^<]{0,80})", xml):
                if formula and not valid_start.match(formula):
                    invalid.append((name, formula))

        self.assertEqual(
            invalid,
            [],
            f"formula tak sah akan dibuang Excel saat dibuka: {invalid[:3]}",
        )

    def test_excel_sheets_are_readable(self):
        """X-2: setiap sheet bertabel harus punya wrap text dan header dibekukan.

        Tanpa ini uraian panjang terpotong dan judul kolom hilang saat menggulir
        -- keluhan keterbacaan yang dilaporkan owner 2026-08-18.
        """
        from io import BytesIO

        from openpyxl import load_workbook

        self.client.force_login(self.owner)
        workbook = load_workbook(BytesIO(self.client.get(self._url("xlsx")).content))

        for name in workbook.sheetnames:
            sheet = workbook[name]
            # "Bertabel" = punya beberapa baris berisi minimal 3 kolom. Lembar
            # pengesahan tidak termasuk: ia blok tanda tangan, bukan tabel data,
            # jadi wrap dan freeze memang tidak relevan di sana.
            dense_rows = sum(
                1
                for row in sheet.iter_rows(max_row=40)
                if sum(1 for cell in row if cell.value not in (None, "")) >= 3
            )
            if dense_rows < 5:
                continue
            with self.subTest(sheet=name):
                wrapped = sum(
                    1
                    for row in sheet.iter_rows(max_row=80)
                    for cell in row
                    if cell.alignment and cell.alignment.wrap_text
                )
                self.assertGreater(wrapped, 0, f"sheet '{name}' tanpa wrap text")
                self.assertIsNotNone(
                    sheet.freeze_panes, f"sheet '{name}' tanpa header dibekukan"
                )
