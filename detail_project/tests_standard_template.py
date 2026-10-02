"""Standar template dokumen export (docs/RENCANA_STANDAR_TEMPLATE_20261002.md).

PDF dokumen perencanaan dirender oleh ``exports/standard``: header/footer
bernomor, huruf DejaVu Sans (glyph ∅), persen id-ID, Rekap Kebutuhan
berkelompok, Paket Perencanaan utuh (cover + daftar isi + 5 dokumen).
"""
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.test import SimpleTestCase, TestCase

from dashboard.models import Project
from detail_project.exports.export_manager import ExportManager
from detail_project.exports.standard import format as fmt
from detail_project.exports.standard import pdf as std
from detail_project.exports.standard.documents import kebutuhan_blocks
from detail_project.exports.table_styles import ROW_MIN_HEIGHT_CM, ROW_MIN_HEIGHT_RINCIAN_CM
from detail_project.models import Klasifikasi, Pekerjaan, SubKlasifikasi
from detail_project.tests_jadwal_monthly_report import pdf_page_texts


class StandardFormattingTests(SimpleTestCase):
    def test_numbers_and_labels_use_id_locale(self):
        self.assertEqual(fmt.num(Decimal('1234567.891')), '1.234.567,89')
        self.assertEqual(fmt.num('73.5', 3), '73,500')
        self.assertEqual(fmt.num(None), '-')
        self.assertEqual(fmt.num('Rp 234/kg'), 'Rp 234/kg')  # teks diteruskan
        self.assertEqual(fmt.id_label('PPN 11.00%'), 'PPN 11,00%')

    def test_all_dash_columns_are_dropped(self):
        headers, rows, widths, aligns = fmt.drop_empty_columns(
            ['No', 'Nama', 'Expression', 'Nilai', 'Satuan'],
            [['1', 'Lebar', '-', '7,00', '-'], ['2', 'Tinggi', '-', '0,20', '']],
            [1, 2, 3, 4, 5], ['c', 'l', 'l', 'r', 'c'], keep=(0, 1, 3))
        self.assertEqual(headers, ['No', 'Nama', 'Nilai'])
        self.assertEqual(rows[0], ['1', 'Lebar', '7,00'])
        self.assertEqual(widths, [1, 2, 4])

    def test_row_minimum_per_table(self):
        regular = std.data_table(['A'], [['x']], [1])
        compact = std.data_table(['A'], [['x']], [1], compact=True)
        self.assertEqual(regular.floor_cm, ROW_MIN_HEIGHT_CM)
        self.assertEqual(compact.floor_cm, ROW_MIN_HEIGHT_RINCIAN_CM)
        self.assertFalse(std.totals_block([('Total', '1')]).enforce_min_row_height)

    def test_kebutuhan_grouped_by_category_with_subtotals(self):
        data = {
            'table_data': {'rows': [
                ['1', 'B-01', 'Semen', 'kg', Decimal('2'), Decimal('10'), Decimal('20')],
                ['2', 'TK-01', 'Pekerja', 'OH', Decimal('1'), Decimal('100'), Decimal('100')],
                ['3', 'PR-01', 'Molen', 'hari', Decimal('1'), Decimal('5'), Decimal('5')],
            ]},
            'row_kategori': ['BHN', 'TK', 'ALT'],
            'footer_rows': [['Total Items', '3']],
        }

        class Ctx:
            project = 'Proyek Uji'
            identity = [('Proyek', 'Proyek Uji')]
            signatures = []

        text = '\n'.join(pdf_page_texts(std.build_pdf(std.flowables(kebutuhan_blocks(data), Ctx()), Ctx(), 'x')))
        self.assertLess(text.index('TENAGA KERJA'), text.index('BAHAN'))
        self.assertLess(text.index('BAHAN'), text.index('ALAT'))
        self.assertIn('Subtotal Tenaga Kerja', text)
        self.assertIn('125,00', text)  # grand total


class StandardPlanningPdfTests(TestCase):
    def setUp(self):
        owner = get_user_model().objects.create_user('std-template', password='x')
        self.project = Project.objects.create(
            owner=owner, nama='Proyek Standar Template', sumber_dana='APBD',
            lokasi_project='Mataram', nama_client='Dinas Uji',
            nama_konsultan_perencana='Perencana Uji', anggaran_owner=Decimal('1000000'),
        )
        klas = Klasifikasi.objects.create(project=self.project, name='Klas A', ordering_index=1)
        sub = SubKlasifikasi.objects.create(project=self.project, klasifikasi=klas, name='Sub A', ordering_index=1)
        Pekerjaan.objects.create(
            project=self.project, sub_klasifikasi=sub, ordering_index=1,
            source_type=Pekerjaan.SOURCE_CUSTOM, snapshot_kode='CUS-001',
            snapshot_uraian='Galian tanah s.d. ∅ 2 cm', snapshot_satuan='m3',
        )

    def _pdf(self, call):
        return call(ExportManager(self.project)).content

    def test_planning_pdfs_use_standard_running_header_and_footer(self):
        exports = {
            'rekap_rab': lambda m: m.export_rekap_rab('pdf'),
            'rincian_ahsp': lambda m: m.export_rincian_ahsp('pdf'),
            'volume_pekerjaan': lambda m: m.export_volume_pekerjaan('pdf'),
            'harga_items': lambda m: m.export_harga_items('pdf'),
            'rekap_kebutuhan': lambda m: m.export_rekap_kebutuhan('pdf'),
        }
        for name, call in exports.items():
            with self.subTest(name=name):
                content = self._pdf(call)
                pages = pdf_page_texts(content)
                self.assertIn(f'Halaman 1 dari {len(pages)}', pages[0])
                self.assertIn('Dashboard-RAB.com', pages[-1])
                self.assertIn('Proyek Standar Template', pages[0])
                self.assertIn(b'DejaVuSans', content)  # S-2: huruf berglyph lengkap tertanam

    def test_rab_font_has_special_symbols_and_percent_is_id(self):
        from reportlab.pdfbase import pdfmetrics

        from detail_project.exports.standard.fonts import font, register

        self.assertTrue(register(), 'berkas DejaVu Sans tidak ditemukan di exports/fonts')
        face = pdfmetrics.getFont(font('Std')).face
        for symbol in '∅½’÷–':
            self.assertIn(ord(symbol), face.charToGlyph, f'glyph {symbol!r} tidak ada')
        text = '\n'.join(pdf_page_texts(self._pdf(lambda m: m.export_rekap_rab('pdf'))))
        self.assertNotRegex(text, r'PPN \d+\.\d+%')

    def test_paket_pdf_is_one_document_with_cover_toc_and_five_documents(self):
        pages = pdf_page_texts(self._pdf(lambda m: m.export_paket_perencanaan('pdf')))
        self.assertIn('DOKUMEN PERENCANAAN', pages[0])
        self.assertNotIn('Halaman', pages[0])  # cover tanpa header/footer
        self.assertIn('DAFTAR ISI', pages[1])
        for title in ('Rencana Anggaran Biaya', 'Rincian Analisa Harga Satuan Pekerjaan',
                      'Volume Pekerjaan', 'Daftar Harga Satuan Dasar', 'Rekap Kebutuhan Material'):
            self.assertIn(title, pages[1])
        self.assertIn(f'Halaman {len(pages)} dari {len(pages)}', pages[-1])


class StandardPlanningWordTests(StandardPlanningPdfTests):
    """Word mengikuti spesifikasi yang sama dengan PDF (S-8)."""

    def _doc(self, call):
        from io import BytesIO

        from docx import Document

        return Document(BytesIO(call(ExportManager(self.project)).content))

    def test_word_has_running_footer_with_page_fields_and_arial(self):
        doc = self._doc(lambda m: m.export_rekap_rab('word'))
        self.assertEqual(doc.styles['Normal'].font.name, 'Arial')
        footer_xml = doc.sections[0].footer._element.xml
        self.assertIn('PAGE', footer_xml)
        self.assertIn('NUMPAGES', footer_xml)
        self.assertIn('Dashboard-RAB.com', footer_xml)
        self.assertIn('Rencana Anggaran Biaya', doc.sections[0].header.paragraphs[0].text)

    def test_word_rincian_rows_use_compact_minimum(self):
        doc = self._doc(lambda m: m.export_rincian_ahsp('word'))
        detail = [t for t in doc.tables if len(t.columns) == 7 and 'Koefisien' in t.rows[0].cells[4].text]
        self.assertTrue(detail, 'tabel rincian tidak ditemukan')
        heights = {round(r.height.cm, 2) for t in detail for r in t.rows if r.height is not None}
        self.assertEqual(heights, {ROW_MIN_HEIGHT_RINCIAN_CM})

    def test_word_paket_has_cover_toc_and_section_per_document(self):
        doc = self._doc(lambda m: m.export_paket_perencanaan('word'))
        self.assertTrue(doc.sections[0].different_first_page_header_footer)  # cover tanpa header
        self.assertIn('TOC', doc.element.body.xml)
        headers = [s.header.paragraphs[0].text for s in doc.sections]
        for title in ('Rencana Anggaran Biaya', 'Analisa Harga Satuan Pekerjaan', 'Volume Pekerjaan',
                      'Daftar Harga Satuan Dasar', 'Rekap Kebutuhan Material'):
            self.assertTrue(any(title in h for h in headers), f'section {title} tidak ada')

    # Tes PDF warisan tidak diulang di kelas ini.
    test_planning_pdfs_use_standard_running_header_and_footer = None
    test_rab_font_has_special_symbols_and_percent_is_id = None
    test_paket_pdf_is_one_document_with_cover_toc_and_five_documents = None
