"""WP-B5 inc-B5e — signature configuration per report (locks DoD)."""
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.test import TestCase

from dashboard.models import Project
from detail_project.exports.signature_config import (
    SignatureLayoutRules,
    SignaturePresets,
)


class SignatureConfigTests(TestCase):
    def setUp(self):
        self.owner = get_user_model().objects.create_user("b5e-owner", password="x")
        self.project = Project.objects.create(
            owner=self.owner,
            nama="Proyek TTD",
            sumber_dana="APBD",
            lokasi_project="Kota",
            nama_client="Dinas PU",
            anggaran_owner=Decimal("1000000"),
            nama_konsultan_perencana="CV Rencana",
            nama_kontraktor="PT Bangun",
            nama_konsultan_pengawas="CV Awas",
        )

    def test_document_presets_map_report_types(self):
        for doc in ["rekap_rab", "rekap_kebutuhan", "volume_pekerjaan", "harga_items", "rincian_ahsp"]:
            with self.subTest(doc=doc):
                self.assertEqual(
                    SignaturePresets.get_roles_for_document(doc),
                    SignaturePresets.PERENCANAAN,
                )
        for doc in ["jadwal_pekerjaan", "jadwal_professional", "laporan_progress"]:
            with self.subTest(doc=doc):
                self.assertEqual(
                    SignaturePresets.get_roles_for_document(doc),
                    SignaturePresets.PELAKSANAAN,
                )

    def test_unknown_document_defaults_to_perencanaan(self):
        self.assertEqual(
            SignaturePresets.get_roles_for_document("anything-else"),
            SignaturePresets.PERENCANAAN,
        )

    def test_build_signatures_perencanaan_owner_and_perencana(self):
        sigs = SignaturePresets.build_signatures(self.project, "PERENCANAAN")
        self.assertEqual([s["label"] for s in sigs], ["Pemilik Proyek", "Konsultan Perencana"])
        self.assertEqual([s["name"] for s in sigs], ["Dinas PU", "CV Rencana"])

    def test_build_signatures_pelaksanaan_three_roles(self):
        sigs = SignaturePresets.build_signatures(self.project, "PELAKSANAAN")
        self.assertEqual(
            [s["label"] for s in sigs],
            ["Pemilik Proyek", "Kontraktor Pelaksana", "Konsultan Pengawas"],
        )
        self.assertEqual([s["name"] for s in sigs], ["Dinas PU", "PT Bangun", "CV Awas"])

    def test_each_role_gets_its_own_instansi(self):
        """Instansi harus milik peran itu sendiri, bukan instansi klien untuk semua."""
        self.project.instansi_client = "Pemkot A"
        self.project.instansi_konsultan_perencana = "CV Rencana Jaya"
        self.project.instansi_kontraktor = "PT Bangun Jaya"
        self.project.instansi_konsultan_pengawas = "CV Awas Jaya"
        self.project.save()

        sigs = SignaturePresets.build_signatures(self.project, "FULL")

        self.assertEqual(
            [s["position"] for s in sigs],
            ["Pemkot A", "CV Rencana Jaya", "PT Bangun Jaya", "CV Awas Jaya"],
        )


class SignatureSheetRenderTests(TestCase):
    """Lembar pengesahan PDF harus benar-benar MENCETAK nama, bukan garis kosong.

    Tes lain di berkas ini berhenti pada ``build_signatures`` -- datanya benar,
    tapi ``PDFExporter._build_signatures`` mengisi baris nama dengan 20 garis
    bawah literal dan tidak pernah memakai ``sig['name']``. Akibatnya lembar
    pengesahan Rekap RAB dan Rincian AHSP terbit kosong meski data konsultan
    lengkap di Dashboard. Tes ini memeriksa tabel yang benar-benar dirender.
    """

    def setUp(self):
        self.owner = get_user_model().objects.create_user("ttd-render", password="x")
        self.project = Project.objects.create(
            owner=self.owner,
            nama="Proyek Pengesahan",
            sumber_dana="APBD",
            lokasi_project="Mataram",
            nama_client="Dinas Peternakan",
            instansi_client="Pemkab Lombok Barat",
            nama_konsultan_perencana="Rozan Fahriady",
            instansi_konsultan_perencana="CV Kaleidoskop",
            anggaran_owner=Decimal("1000000"),
        )

    def _signature_rows(self):
        from reportlab.platypus import Table

        from detail_project.exports.export_manager import ExportManager
        from detail_project.exports.pdf_exporter import PDFExporter

        elements = PDFExporter(ExportManager(self.project)._create_config())._build_signatures()

        def _tables(obj):
            if isinstance(obj, Table):
                yield obj
            for attr in ("_content", "_flowables"):
                for child in getattr(obj, attr, []) or []:
                    yield from _tables(child)

        rows = []
        for table in _tables(elements[0]):
            rows.extend(table._cellvalues)
        return rows

    @staticmethod
    def _plain(cell):
        """Teks sel tanpa markup.

        Sel nama berupa Paragraph dengan markup ``<u>`` (garis bawah selebar
        teks), sisanya string biasa.
        """
        import re

        return re.sub(r"<[^>]+>", "", str(getattr(cell, "text", cell)))

    def _flat(self):
        return [self._plain(c) for row in self._signature_rows() for c in row]

    def test_names_and_instansi_are_printed(self):
        flat = self._flat()

        for expected in (
            "Dinas Peternakan",
            "Rozan Fahriady",
            "Pemkab Lombok Barat",
            "CV Kaleidoskop",
        ):
            with self.subTest(expected=expected):
                self.assertIn(
                    expected,
                    flat,
                    f"'{expected}' tidak tercetak di lembar pengesahan PDF.",
                )

    def test_no_connective_or_role_label_is_printed(self):
        """R-37 (owner 2026-09-27): susunan hanya Instansi / Nama / Jabatan.

        Kata penghubung dan sebutan peran tidak dicetak, termasuk sebutan
        kustom ``sebutan_client``.
        """
        self.project.sebutan_client = "Pejabat Pembuat Komitmen Dinas Peternakan"
        self.project.save()

        flat = self._flat()
        for removed in (
            "Mengetahui,", "Dibuat oleh,", "Diperiksa oleh,",
            "Pemilik Proyek", "Konsultan Perencana",
            "Pejabat Pembuat Komitmen Dinas Peternakan",
        ):
            with self.subTest(removed=removed):
                self.assertNotIn(removed, flat)

    def test_column_order_is_instansi_name_ket2_jabatan(self):
        """Per pihak: INSTANSI, (ruang TTD), NAMA, lalu ket_client2, lalu jabatan."""
        self.project.jabatan_client = "PPK Konstruksi"
        self.project.ket_client2 = "NIP 19700101 199003 1 001"
        self.project.save()

        rows = self._signature_rows()
        owner_col = [self._plain(r[0]) for r in rows]
        self.assertEqual(
            [v for v in owner_col if v],
            ["Pemkab Lombok Barat", "Dinas Peternakan", "NIP 19700101 199003 1 001", "PPK Konstruksi"],
        )
        # Pihak tanpa jabatan (perencana) tidak mendapat baris keterangan isian.
        perencana_col = [self._plain(r[1]) for r in rows]
        self.assertEqual([v for v in perencana_col if v], ["CV Kaleidoskop", "Rozan Fahriady"])

    def test_signature_table_fits_the_printable_width(self):
        """Lebar tabel dulu 250mm mati, melebihi area cetak A4 portrait 190mm."""
        from reportlab.lib.units import mm
        from reportlab.platypus import Table

        from detail_project.exports.export_manager import ExportManager
        from detail_project.exports.pdf_exporter import PDFExporter

        config = ExportManager(self.project)._create_config()
        exporter = PDFExporter(config)
        printable_mm = exporter._get_page_width_mm() - (
            config.margin_left + config.margin_right
        )

        def _tables(obj):
            if isinstance(obj, Table):
                yield obj
            for attr in ("_content", "_flowables"):
                for child in getattr(obj, attr, []) or []:
                    yield from _tables(child)

        for table in _tables(exporter._build_signatures()[0]):
            total_mm = sum(table._colWidths) / mm
            self.assertLessEqual(
                round(total_mm, 1),
                round(printable_mm, 1),
                f"tabel tanda tangan {total_mm:.0f}mm melebihi area cetak "
                f"{printable_mm:.0f}mm -- kolom kanan akan terpotong.",
            )

    def test_name_is_underlined_and_no_separate_line_row(self):
        """Keputusan owner: garis di ATAS nama dihapus, namanya yang digarisbawahi.

        Garis bawah dibungkus Paragraph (markup ``<u>``) supaya panjangnya
        mengikuti teks, bukan selebar kolom seperti LINEBELOW.
        """
        rows = self._signature_rows()

        underscore_rows = [
            r for r in rows
            if any(self._plain(c).strip("_") == "" and self._plain(c) for c in r)
        ]
        self.assertEqual(underscore_rows, [], "baris garis bawah terpisah masih ada")

        name_cells = [
            c for r in rows for c in r
            if "Rozan Fahriady" in self._plain(c)
        ]
        self.assertTrue(name_cells, "nama tidak tercetak")
        self.assertIn("<u>", str(getattr(name_cells[0], "text", "")),
                      "nama tidak digarisbawahi")


class SignatureWordExcelRenderTests(TestCase):
    """R-37 di Word (blok baku) dan Excel (sheet Pengesahan Volume)."""

    def setUp(self):
        self.owner = get_user_model().objects.create_user("ttd-word-excel", password="x")
        self.project = Project.objects.create(
            owner=self.owner,
            nama="Proyek Pengesahan",
            sumber_dana="APBD",
            lokasi_project="Mataram",
            nama_client="Dinas Peternakan",
            instansi_client="Pemkab Lombok Barat",
            nama_konsultan_perencana="Rozan Fahriady",
            instansi_konsultan_perencana="CV Kaleidoskop",
            anggaran_owner=Decimal("1000000"),
        )

    def test_word_block_is_instansi_space_name_details(self):
        from io import BytesIO

        from docx import Document

        from detail_project.exports.export_manager import ExportManager

        self.project.jabatan_client = "PPK Konstruksi"
        self.project.ket_client2 = "NIP 1970"
        self.project.save()
        doc = Document(BytesIO(ExportManager(self.project).export_rekap_rab("word").content))
        table = next(t for t in doc.tables if t.rows[0].cells[0].text == "Pemkab Lombok Barat")
        rows = [[c.text for c in r.cells] for r in table.rows]

        self.assertEqual(rows[0], ["Pemkab Lombok Barat", "CV Kaleidoskop"])
        self.assertEqual(rows[1], ["", ""])  # ruang tanda tangan basah
        self.assertEqual(rows[2], ["Dinas Peternakan", "Rozan Fahriady"])
        self.assertEqual([r[0] for r in rows[3:]], ["NIP 1970", "PPK Konstruksi"])
        self.assertTrue(all(run.underline for run in table.cell(2, 0).paragraphs[0].runs))
        flat = " ".join(" ".join(r) for r in rows)
        for stale in ("Mengetahui,", "Dibuat oleh,", "Pemilik Proyek", "Konsultan Perencana"):
            self.assertNotIn(stale, flat)

    def test_excel_volume_signature_sheet_uses_project_data(self):
        from io import BytesIO

        from openpyxl import load_workbook

        from detail_project.exports.export_manager import ExportManager

        wb = load_workbook(BytesIO(ExportManager(self.project).export_volume_pekerjaan("xlsx").content))
        values = [str(c.value) for row in wb["Pengesahan"].iter_rows() for c in row if c.value]

        self.assertEqual(
            values,
            ["LEMBAR PENGESAHAN", "Pemkab Lombok Barat", "CV Kaleidoskop", "Dinas Peternakan", "Rozan Fahriady"],
        )


class SignatureCoverageTests(TestCase):
    """Semua dokumen tahap perencanaan harus menerbitkan lembar pengesahan.

    Sebelumnya cakupannya bocor tanpa disadari: hanya Rekap RAB dan Rekap
    Kebutuhan yang memakai blok bersama. Rincian AHSP punya tabel hardcode
    sendiri (label tetap + "(______)"), sementara Volume dan Harga Items tidak
    menerbitkan lembar pengesahan sama sekali di PDF/Word -- hanya di Excel,
    dengan nama berupa titik-titik.

    Tes ini memanggil setiap export sungguhan dan memastikan blok bersama
    benar-benar terpakai, sehingga penambahan dokumen baru tidak diam-diam
    kembali memakai tabel hardcode.
    """

    def setUp(self):
        self.owner = get_user_model().objects.create_user("ttd-coverage", password="x")
        self.project = Project.objects.create(
            owner=self.owner,
            nama="Proyek Cakupan",
            sumber_dana="APBD",
            lokasi_project="Mataram",
            nama_client="Dinas Peternakan",
            nama_konsultan_perencana="Rozan Fahriady",
            anggaran_owner=Decimal("1000000"),
        )

    def test_every_perencanaan_document_emits_the_shared_block(self):
        from detail_project.exports.export_manager import ExportManager
        from detail_project.exports.pdf_exporter import PDFExporter

        original = PDFExporter._build_signatures
        called = set()

        def spy(exporter_self):
            called.add(spy.current)
            return original(exporter_self)

        exports = {
            "rekap_rab": lambda m: m.export_rekap_rab("pdf"),
            "rincian_ahsp": lambda m: m.export_rincian_ahsp("pdf"),
            "volume_pekerjaan": lambda m: m.export_volume_pekerjaan("pdf"),
            "harga_items": lambda m: m.export_harga_items("pdf"),
            "rekap_kebutuhan": lambda m: m.export_rekap_kebutuhan("pdf"),
        }

        PDFExporter._build_signatures = spy
        try:
            for name, call in exports.items():
                spy.current = name
                call(ExportManager(self.project))
        finally:
            PDFExporter._build_signatures = original

        missing = sorted(set(exports) - called)
        self.assertEqual(
            missing,
            [],
            f"dokumen ini tidak menerbitkan lembar pengesahan bersama: {missing}",
        )
