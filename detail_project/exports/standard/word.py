"""Renderer Word standar dokumen perencanaan (python-docx).

Menerjemahkan blok netral yang sama dengan PDF (documents.py) sehingga isi dan
susunan Word = PDF (S-8). Huruf Arial; warna, ukuran, garis, tinggi baris
minimum, dan pengesahan mengikuti token yang sama (tokens.py).
"""
from __future__ import annotations

from io import BytesIO
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Mm, Pt, RGBColor

from . import tokens as T
from ..table_styles import ROW_MIN_HEIGHT_CM, ROW_MIN_HEIGHT_RINCIAN_CM

FONT = 'Arial'
_ALIGN = {'l': WD_ALIGN_PARAGRAPH.LEFT, 'c': WD_ALIGN_PARAGRAPH.CENTER, 'r': WD_ALIGN_PARAGRAPH.RIGHT}


def rgb(color: str) -> RGBColor:
    return RGBColor.from_string(T.word_hex(color))


# ------------------------------------------------------------- XML kecil ----
def _shade(cell, color: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn('w:shd')):
        tc_pr.remove(old)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), T.word_hex(color))
    tc_pr.append(shd)


def _cell_borders(cell, edges: Dict[str, Tuple[str, str]]):
    """edges: {'top': (warna, ukuran_1/8pt), ...}; ukuran 'nil' = tanpa garis."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge, (color, size) in edges.items():
        el = borders.find(qn(f'w:{edge}'))
        if el is None:
            el = OxmlElement(f'w:{edge}')
            borders.append(el)
        if size == 'nil':
            el.set(qn('w:val'), 'nil')
        else:
            el.set(qn('w:val'), 'single')
            el.set(qn('w:sz'), size)
            el.set(qn('w:color'), T.word_hex(color))


_TBLPR_ORDER = ('tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize', 'tblStyleColBandSize',
                'tblW', 'jc', 'tblCellSpacing', 'tblInd', 'tblBorders', 'shd', 'tblLayout', 'tblCellMar',
                'tblLook', 'tblCaption', 'tblDescription')


def _tblpr_insert(tbl_pr, element):
    """Sisipkan anak tblPr sesuai urutan skema OOXML (Word mengabaikan yang salah urutan)."""
    name = element.tag.split('}')[1]
    after = _TBLPR_ORDER[_TBLPR_ORDER.index(name) + 1:]
    for child in tbl_pr:
        if child.tag.split('}')[1] in after:
            child.addprevious(element)
            return
    tbl_pr.append(element)


def _table_borders(table, color: Optional[str], size: str = '4'):
    tbl_pr = table._tbl.tblPr
    for old in tbl_pr.findall(qn('w:tblBorders')):
        tbl_pr.remove(old)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        if color is None:
            el.set(qn('w:val'), 'nil')
        else:
            el.set(qn('w:val'), 'single')
            el.set(qn('w:sz'), size)
            el.set(qn('w:color'), T.word_hex(color))
        borders.append(el)
    _tblpr_insert(tbl_pr, borders)


def _table_cell_margins(table, top_mm: float, side_mm: float):
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement('w:tblCellMar')
    for edge, value in (('top', top_mm), ('bottom', top_mm), ('left', side_mm), ('right', side_mm)):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:w'), str(int(Mm(value).twips)))
        el.set(qn('w:type'), 'dxa')
        mar.append(el)
    for old in tbl_pr.findall(qn('w:tblCellMar')):
        tbl_pr.remove(old)
    _tblpr_insert(tbl_pr, mar)


def _widths(table, widths: Sequence[int]):
    table.autofit = False
    for grid_col, width in zip(table._tbl.tblGrid.findall(qn('w:gridCol')), widths):
        grid_col.set(qn('w:w'), str(int(Emu(int(width)).twips)))
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx < len(row.cells):
                row.cells[idx].width = Emu(int(width))


def _para_border(paragraph, edge: str, color: str, size: str = '4'):
    p_pr = paragraph._p.get_or_add_pPr()
    bdr = p_pr.find(qn('w:pBdr'))
    if bdr is None:
        bdr = OxmlElement('w:pBdr')
        p_pr.append(bdr)
    el = OxmlElement(f'w:{edge}')
    el.set(qn('w:val'), 'single')
    el.set(qn('w:sz'), size)
    el.set(qn('w:space'), '1')
    el.set(qn('w:color'), T.word_hex(color))
    bdr.append(el)


def _field(paragraph, code: str, size: float, color: str):
    run = paragraph.add_run()
    for tag, attr in (('w:fldChar', 'begin'), ('w:instrText', None), ('w:fldChar', 'separate'),
                      ('w:t', None), ('w:fldChar', 'end')):
        el = OxmlElement(tag)
        if tag == 'w:fldChar':
            el.set(qn('w:fldCharType'), attr)
        elif tag == 'w:instrText':
            el.set(qn('xml:space'), 'preserve')
            el.text = f' {code} '
        else:
            el.text = '1'
        run._r.append(el)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)


def _text(paragraph, text, size=T.SIZE_BODY, bold=False, italic=False, color=T.TEXT,
          align=None, underline=False, leading=None):
    if align is not None:
        paragraph.alignment = align
    if leading:
        # Spasi tepat (tabel ringkas Rincian): 'single' Word untuk huruf kecil
        # lebih tinggi dari hurufnya sehingga baris melewati 0,4 cm.
        paragraph.paragraph_format.line_spacing = Pt(leading)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run('' if text is None else str(text))
    run.font.size = Pt(size)
    # Tanda paragraf ikut ukuran teks; bila tidak, ukuran Normal menambah tinggi baris.
    p_pr = paragraph._p.get_or_add_pPr()
    mark = p_pr.find(qn('w:rPr'))
    if mark is None:
        mark = OxmlElement('w:rPr')
        p_pr.append(mark)
    for old in mark.findall(qn('w:sz')):
        mark.remove(old)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), str(int(round(size * 2))))
    mark.append(sz)
    run.bold, run.italic, run.underline = bold, italic, underline
    run.font.color.rgb = rgb(color)
    return run


def _row_height(row, cm: float, exact: bool = False):
    row.height = Cm(cm)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY if exact else WD_ROW_HEIGHT_RULE.AT_LEAST


def _repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), 'true')
    tr_pr.append(el)


def _cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement('w:cantSplit'))


# ----------------------------------------------------------------- kelas ----
class WordBuilder:
    """Satu dokumen Word standar."""

    def __init__(self, config):
        from ...export_config import build_identity_rows

        self.config = config
        self.project = getattr(config, 'project_name', '') or ''
        self.identity = [(row[0], row[2]) for row in build_identity_rows(config)]
        sig_cfg = getattr(config, 'signature_config', None)
        sigs = []
        if sig_cfg is not None and getattr(sig_cfg, 'enabled', False):
            sigs = sig_cfg.custom_signatures or sig_cfg.signatures or []
        self.signatures = list(sigs)
        self.doc = Document()
        self._setup()

    # -- dokumen & header/footer ------------------------------------------
    def _setup(self):
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin = sec.right_margin = Mm(T.PAGE_MARGIN_LR)
        sec.top_margin, sec.bottom_margin = Mm(T.PAGE_MARGIN_TOP), Mm(T.PAGE_MARGIN_BOTTOM)
        sec.header_distance, sec.footer_distance = Mm(8), Mm(7)
        self.width = int(sec.page_width - sec.left_margin - sec.right_margin)
        normal = self.doc.styles['Normal']
        normal.font.name = FONT
        normal.font.size = Pt(T.SIZE_BODY)
        normal.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
        normal.paragraph_format.space_after = Pt(0)
        normal.paragraph_format.line_spacing = 1.0  # bawaan Word ±1,15 membuat baris Rincian > 0,4 cm
        # Template bawaan python-docx memasang document grid; Word lalu
        # menyelaraskan baris ke kisi sehingga baris Rincian ±5,4 mm, bukan 0,4 cm.
        for grid in sec._sectPr.findall(qn('w:docGrid')):
            sec._sectPr.remove(grid)
        p_pr = normal.element.get_or_add_pPr()
        snap = OxmlElement('w:snapToGrid')
        snap.set(qn('w:val'), '0')
        p_pr.insert(0, snap)
        for name in ('Heading 1', 'Heading 2'):
            style = self.doc.styles[name]
            rf = style.element.rPr.rFonts
            for attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
                rf.attrib.pop(qn(attr), None)
            rf.set(qn('w:ascii'), FONT)
            rf.set(qn('w:hAnsi'), FONT)
        upd = OxmlElement('w:updateFields')
        upd.set(qn('w:val'), 'true')
        self.doc.settings.element.append(upd)

    def running(self, section, doc_name: str, blank_first_page: bool = False):
        """Header (proyek | dokumen) + footer 'Halaman x dari y' | brand."""
        section.different_first_page_header_footer = blank_first_page
        for part in (section.header, section.footer):
            part.is_linked_to_previous = False
        hp = section.header.paragraphs[0]
        hp.text = ''
        tabs = hp.paragraph_format.tab_stops
        for stop in (Mm(82.55), Mm(165.1)):  # tab bawaan gaya Header
            tabs.add_tab_stop(stop, WD_TAB_ALIGNMENT.CLEAR)
        tabs.add_tab_stop(Emu(self.width), WD_TAB_ALIGNMENT.RIGHT)
        name = self.project if len(self.project) <= 72 else self.project[:69].rsplit(' ', 1)[0] + '…'
        _text(hp, name, T.SIZE_RUNNING, color=T.TEXT2)
        _text(hp, f'\t{doc_name}', T.SIZE_RUNNING, color=T.TEXT2)
        _para_border(hp, 'bottom', T.GRID)
        fp = section.footer.paragraphs[0]
        fp.text = ''
        tabs = fp.paragraph_format.tab_stops
        for stop in (Mm(82.55), Mm(165.1)):
            tabs.add_tab_stop(stop, WD_TAB_ALIGNMENT.CLEAR)
        tabs.add_tab_stop(Emu(self.width // 2), WD_TAB_ALIGNMENT.CENTER)
        tabs.add_tab_stop(Emu(self.width), WD_TAB_ALIGNMENT.RIGHT)
        _para_border(fp, 'top', T.GRID)
        _text(fp, '\tHalaman ', T.SIZE_RUNNING, color=T.MUTED)
        _field(fp, 'PAGE', T.SIZE_RUNNING, T.MUTED)
        _text(fp, ' dari ', T.SIZE_RUNNING, color=T.MUTED)
        _field(fp, 'NUMPAGES', T.SIZE_RUNNING, T.MUTED)
        _text(fp, f'\t{T.BRAND}', T.SIZE_RUNNING, color=T.MUTED)
        if blank_first_page:
            for part in (section.first_page_header, section.first_page_footer):
                part.is_linked_to_previous = False
                part.paragraphs[0].text = ''

    def new_section(self, doc_name: str):
        section = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self.running(section, doc_name)
        return section

    # -- komponen ---------------------------------------------------------
    def paragraph(self, text='', **kw):
        p = self.doc.add_paragraph()
        if text:
            _text(p, text, **kw)
        return p

    def spacer(self, mm_value: float):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(max(1, mm_value * 2.835))
        return p

    def doc_header(self, title: str, subtitle: Optional[str] = None):
        p = self.paragraph(title, size=T.SIZE_DOC_TITLE, bold=True, color=T.NAVY_DARK,
                           align=WD_ALIGN_PARAGRAPH.CENTER)
        p.style = self.doc.styles['Heading 1']  # navigasi Word; tampilan ditimpa run
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.keep_with_next = True
        if subtitle:
            self.paragraph(subtitle, size=T.SIZE_SUBTITLE, color=T.TEXT2, align=WD_ALIGN_PARAGRAPH.CENTER)
        rule = self.doc.add_paragraph()
        indent = Emu(int((self.width - Mm(T.TITLE_RULE_WIDTH)) / 2))
        rule.paragraph_format.left_indent = indent
        rule.paragraph_format.right_indent = indent
        rule.paragraph_format.space_before = Pt(2)
        rule.paragraph_format.space_after = Pt(9)
        rule.paragraph_format.line_spacing = Pt(2)
        _para_border(rule, 'bottom', T.ACCENT, '6')
        self.identity_panel()
        self.spacer(4)

    def identity_panel(self):
        identity = self.identity
        half = (len(identity) + 1) // 2
        left, right = identity[:half], identity[half:]
        t = self.doc.add_table(rows=half, cols=7)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _table_borders(t, None)
        _table_cell_margins(t, 0.9, 1.4)
        w = self.width
        _widths(t, [Mm(27), Mm(3), w / 2 - Mm(33), Mm(6), Mm(28), Mm(3), w / 2 - Mm(34)])
        for i in range(half):
            for base, rows in ((0, left), (4, right)):
                label, value = rows[i] if i < len(rows) else ('', '')
                cells = t.rows[i].cells
                _text(cells[base].paragraphs[0], label, bold=True, color=T.TEXT2)
                _text(cells[base + 1].paragraphs[0], ':' if label else '', bold=True, color=T.TEXT2)
                _text(cells[base + 2].paragraphs[0], value)
                for off in range(3):
                    edges = {}
                    if i == 0:
                        edges['top'] = (T.GRID, '4')
                    if i == half - 1:
                        edges['bottom'] = (T.GRID, '4')
                    if off == 0:
                        edges['left'] = (T.GRID, '4')
                    if off == 2:
                        edges['right'] = (T.GRID, '4')
                    if edges:
                        _cell_borders(cells[base + off], edges)
        return t

    def data_table(self, spec: Dict):
        headers, rows, kinds, aligns = spec['headers'], spec['rows'], spec['kinds'], spec['aligns']
        compact, span_to = spec['compact'], spec['span_to']
        ncol = len(headers)
        total = float(sum(spec['widths']))
        widths = [int(self.width * w / total) for w in spec['widths']]
        size = T.RINCIAN_SIZE_BODY if compact else T.SIZE_BODY
        lead = T.RINCIAN_LEADING if compact else None
        min_cm = ROW_MIN_HEIGHT_RINCIAN_CM if compact else ROW_MIN_HEIGHT_CM
        t = self.doc.add_table(rows=0, cols=ncol)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _table_borders(t, T.GRID, '4')
        # Margin atas/bawah 0: di Word tinggi minimum baris (atLeast) TIDAK
        # mencakup margin sel, jadi margin vertikal akan menambah tinggi
        # (0,4 cm + 2x0,63 mm = 5,3 mm). Ruang vertikal = tinggi minimum baris.
        _table_cell_margins(t, 0, T.RINCIAN_PAD_H_MM if compact else T.CELL_PAD_H)
        if spec['header']:
            row = t.add_row()
            _row_height(row, min_cm)
            _repeat_header(row)
            for i, h in enumerate(headers):
                cell = row.cells[i]
                _text(cell.paragraphs[0], h, size, bold=True, color=T.WHITE, align=WD_ALIGN_PARAGRAPH.CENTER,
                      leading=lead)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                _shade(cell, T.NAVY_DARK)
        for values, kind in zip(rows, kinds):
            values = list(values) + [''] * (ncol - len(values))
            row = t.add_row()
            _row_height(row, min_cm)
            _cant_split(row)
            cells = row.cells
            if kind in ('klas', 'sub'):
                last = ncol - 1 if span_to is None else span_to
                label = next((c for c in values if c not in (None, '')), '')
                merged = cells[0].merge(cells[last]) if last > 0 else cells[0]
                fsize = (7 if compact else T.SIZE_CATEGORY) if kind == 'klas' else \
                    (T.RINCIAN_SIZE_GROUP if compact else T.SIZE_SUBCATEGORY)
                _text(merged.paragraphs[0], label, fsize, bold=True, italic=(kind == 'sub'), leading=lead)
                if span_to is not None:
                    for c in range(last + 1, ncol):
                        if values[c] not in (None, ''):
                            _text(cells[c].paragraphs[0], values[c], fsize, bold=True, italic=(kind == 'sub'),
                                  align=WD_ALIGN_PARAGRAPH.RIGHT)
                if kind == 'klas':
                    for cell in row.cells:
                        _shade(cell, T.KLAS_BG)
            elif kind in ('total', 'grand'):
                label = next((c for c in values[:-1] if c not in (None, '')), '')
                merged = cells[0].merge(cells[ncol - 2]) if ncol > 2 else cells[0]
                _text(merged.paragraphs[0], label, size, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, leading=lead)
                _text(cells[ncol - 1].paragraphs[0], values[-1], size, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT,
                      leading=lead)
                if kind == 'grand':
                    for cell in row.cells:
                        _shade(cell, T.GRAND_BG)
            else:
                for i in range(ncol):
                    _text(cells[i].paragraphs[0], values[i], size, align=_ALIGN[aligns[i]], leading=lead)
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _widths(t, widths)
        return t

    def totals(self, pairs, emphasize_last: bool = True):
        t = self.doc.add_table(rows=len(pairs), cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.RIGHT
        _table_borders(t, None)
        _table_cell_margins(t, 0.9, 1.2)
        _widths(t, [Mm(T.TOTAL_BLOCK_WIDTH[0]), Mm(T.TOTAL_BLOCK_WIDTH[1])])
        for i, (label, value) in enumerate(pairs):
            cells = t.rows[i].cells
            _text(cells[0].paragraphs[0], label, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
            _text(cells[1].paragraphs[0], value, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
            if i == 0:
                for cell in cells:
                    _cell_borders(cell, {'top': (T.ACCENT, '6')})
        if emphasize_last and len(pairs) > 1:
            for cell in t.rows[-1].cells:
                _cell_borders(cell, {'top': (T.NAVY_DARK, '10')})
                _shade(cell, T.GRAND_BG)
        return t

    def signature(self):
        from ..signature_config import SIGNATURE_SPACE_MM

        sigs = self.signatures
        if not sigs:
            return
        p = self.paragraph('LEMBAR PENGESAHAN', size=9, bold=True)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.keep_with_next = True
        details = [list(s.get('details') or []) for s in sigs]
        max_det = max((len(d) for d in details), default=0)
        t = self.doc.add_table(rows=3 + max_det, cols=len(sigs))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _table_borders(t, None)
        _widths(t, [self.width // len(sigs)] * len(sigs))
        _row_height(t.rows[1], SIGNATURE_SPACE_MM / 10, exact=True)  # R-23
        for col, sig in enumerate(sigs):
            center = WD_ALIGN_PARAGRAPH.CENTER
            _text(t.rows[0].cells[col].paragraphs[0], sig.get('instansi') or sig.get('position') or '',
                  T.SIZE_SIGNATURE, bold=True, align=center)
            _text(t.rows[2].cells[col].paragraphs[0], sig.get('name') or '', T.SIZE_SIGNATURE,
                  bold=True, underline=True, align=center)
            for i in range(max_det):
                text = details[col][i] if i < len(details[col]) else ''
                _text(t.rows[3 + i].cells[col].paragraphs[0], text, 7.5, color=T.TEXT2, align=center)
            t.rows[0].cells[col].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
        for row in t.rows:
            _cant_split(row)
        return t

    # -- blok -------------------------------------------------------------
    def blocks(self, blocks: list):
        for block in blocks:
            kind = block[0]
            if kind == 'header':
                self.doc_header(block[1], block[2])
            elif kind == 'section':
                p = self.paragraph(block[1], size=T.SIZE_SECTION, bold=True, color=T.ACCENT)
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.keep_with_next = True
            elif kind == 'heading':
                p = self.paragraph(block[1], size=T.RINCIAN_SIZE_HEADING, bold=True, color=T.ACCENT)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.keep_with_next = True
            elif kind == 'table':
                self.data_table(block[1])
            elif kind == 'totals':
                self.totals(block[1], block[2])
            elif kind == 'note':
                p = self.paragraph(block[1], size=T.SIZE_NOTE, italic=True, color=T.MUTED)
                p.paragraph_format.space_before = Pt(3)
            elif kind == 'spacer':
                self.spacer(block[1])
            elif kind == 'signature':
                self.signature()
            elif kind == 'keep':
                start = len(self.doc.element.body)
                self.blocks(block[1])
                self._keep_together(start)
            elif kind == 'pagebreak':
                self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            else:  # pragma: no cover
                raise ValueError(f'blok tidak dikenal: {kind}')

    def _keep_together(self, start: int):
        """R-13: semua paragraf blok (termasuk di dalam sel) menempel ke berikutnya."""
        body = self.doc.element.body
        new = [el for el in list(body)[start:] if el.tag != qn('w:sectPr')]
        paras = []
        for el in new:
            paras += el.iter(qn('w:p'))
        for p in paras[:-1]:
            p_pr = p.get_or_add_pPr()
            if p_pr.find(qn('w:keepNext')) is None:
                keep = OxmlElement('w:keepNext')
                mark = p_pr.find(qn('w:rPr'))  # skema: keepNext harus sebelum rPr
                if mark is not None:
                    mark.addprevious(keep)
                else:
                    p_pr.append(keep)

    # -- cover & daftar isi (paket) ---------------------------------------
    def cover(self, kind: str, contents: str):
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl_pr = t._tbl.tblPr
        borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right'):
            el = OxmlElement(f'w:{edge}')
            el.set(qn('w:val'), 'single')
            el.set(qn('w:sz'), '12')
            el.set(qn('w:color'), T.word_hex(T.ACCENT))
            borders.append(el)
        _tblpr_insert(tbl_pr, borders)
        _widths(t, [self.width])
        _row_height(t.rows[0], 25.0, exact=True)
        cell = t.rows[0].cells[0]
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Mm(55)
        _text(p, kind, 22, bold=True, color=T.NAVY_DARK, align=WD_ALIGN_PARAGRAPH.CENTER)
        p.paragraph_format.space_before = Mm(55)
        rule = cell.add_paragraph()
        rule.paragraph_format.left_indent = rule.paragraph_format.right_indent = Mm(55)
        rule.paragraph_format.space_after = Mm(6)
        _para_border(rule, 'bottom', T.ACCENT, '8')
        for text, size, bold, color in ((self.project, 15, True, T.NAVY_DARK), (contents, 9, False, T.TEXT2)):
            p = cell.add_paragraph()
            p.paragraph_format.left_indent = p.paragraph_format.right_indent = Mm(12)
            p.paragraph_format.space_after = Mm(5)
            _text(p, text, size, bold=bold, color=color, align=WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.space_after = Mm(5)
        spacer = cell.add_paragraph()
        spacer.paragraph_format.space_before = Mm(8)
        info = cell.add_table(rows=0, cols=3)
        info.alignment = WD_TABLE_ALIGNMENT.CENTER
        _table_borders(info, None)
        _table_cell_margins(info, 1.4, 1.0)
        for label, value in self.identity:
            if label.lower() == 'proyek':
                continue
            cells = info.add_row().cells
            _text(cells[0].paragraphs[0], label, bold=True, color=T.TEXT2)
            _text(cells[1].paragraphs[0], ':', bold=True, color=T.TEXT2)
            _text(cells[2].paragraphs[0], value)
        _widths(info, [Mm(38), Mm(4), Mm(100)])

    def toc(self):
        p = self.paragraph('DAFTAR ISI', size=T.SIZE_DOC_TITLE, bold=True, color=T.NAVY_DARK,
                           align=WD_ALIGN_PARAGRAPH.CENTER)
        p.paragraph_format.space_after = Pt(12)
        toc = self.doc.add_paragraph()
        _field(toc, 'TOC \\o "1-1" \\h \\z \\u', 9.5, T.TEXT)

    def save(self) -> bytes:
        buf = BytesIO()
        self.doc.core_properties.author = 'Dashboard-RAB'
        self.doc.save(buf)
        return buf.getvalue()


# ----------------------------------------------------------------- API ----
def build_word(blocks: list, config, doc_name: str) -> bytes:
    builder = WordBuilder(config)
    builder.running(builder.doc.sections[0], doc_name)
    builder.blocks(blocks)
    return builder.save()


def build_word_package(documents: Iterable[Tuple[list, str]], config, contents: str) -> bytes:
    """Paket: cover (tanpa header/footer) + daftar isi + satu section per dokumen."""
    builder = WordBuilder(config)
    builder.running(builder.doc.sections[0], 'Dokumen Perencanaan', blank_first_page=True)
    builder.cover('DOKUMEN PERENCANAAN', contents)
    builder.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    builder.toc()
    for blocks, doc_name in documents:
        builder.new_section(doc_name)
        builder.blocks(blocks)
    return builder.save()
