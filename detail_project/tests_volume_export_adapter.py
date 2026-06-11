import re
from io import BytesIO
from unittest import skipUnless

from django.contrib.auth import get_user_model
from django.test import TestCase

from dashboard.models import Project
from detail_project.exports.excel_exporter import OPENPYXL_AVAILABLE
from detail_project.exports.export_manager import ExportManager
from detail_project.exports.volume_pekerjaan_adapter import VolumePekerjaanAdapter
from detail_project.models import (
    Klasifikasi,
    Pekerjaan,
    ProjectComputedParameter,
    ProjectParameter,
    SubKlasifikasi,
    VolumeFormulaState,
    VolumePekerjaan,
)
from detail_project.services import DeepCopyService

if OPENPYXL_AVAILABLE:
    from openpyxl import load_workbook


class VolumeExportAdapterHardeningTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.owner = user_model.objects.create_user(
            username="volume_export_adapter_user",
            email="volume-export-adapter@example.com",
            password="Secret123!",
        )
        self.project = Project.objects.create(owner=self.owner, nama="Volume Export Adapter")

    def _create_one_pekerjaan(self):
        klas = Klasifikasi.objects.create(project=self.project, name="Klasifikasi A")
        sub = SubKlasifikasi.objects.create(project=self.project, klasifikasi=klas, name="Sub A")
        pekerjaan = Pekerjaan.objects.create(
            project=self.project,
            sub_klasifikasi=sub,
            source_type=Pekerjaan.SOURCE_CUSTOM,
            snapshot_kode="CUST-0001",
            snapshot_uraian="Pekerjaan Uji",
            snapshot_satuan="m3",
        )
        VolumePekerjaan.objects.create(project=self.project, pekerjaan=pekerjaan, quantity="10")
        return pekerjaan

    def test_adapter_normalizes_parameter_payload_and_keeps_label_formula(self):
        ProjectParameter.objects.create(
            project=self.project,
            name="bp_1",
            label="Panjang Dinding",
            value="7",
        )
        pekerjaan = self._create_one_pekerjaan()
        VolumeFormulaState.objects.create(
            project=self.project,
            pekerjaan=pekerjaan,
            raw="=bp_1 * 2",
            is_fx=True,
        )

        adapter = VolumePekerjaanAdapter(
            self.project,
            parameters={"BP_1": {"value": 12}},
        )
        data = adapter.get_export_data()

        param_rows = data["pages"][1]["table_data"]["rows"]
        headers = data["pages"][1]["table_data"]["headers"]
        self.assertEqual(headers, ["No", "Nama Parameter", "Nilai", "Satuan"])

        bp_rows = [row for row in param_rows if len(row) >= 4 and row[1] == "Panjang Dinding"]
        self.assertEqual(len(bp_rows), 1)
        self.assertEqual(bp_rows[0][2], "12")
        self.assertEqual(bp_rows[0][3], "-")
        self.assertEqual(data["pages"][1]["table_data"]["param_codes"], ["bp_1"])

        volume_rows = data["pages"][0]["table_data"]["rows"]
        item_rows = [row for row in volume_rows if len(row) >= 3 and row[0] == "1"]
        self.assertEqual(len(item_rows), 1)
        self.assertIn("Panjang Dinding", item_rows[0][2])
        self.assertNotIn("bp_1", item_rows[0][2].lower())

        row_formulas = data["pages"][0]["row_formulas"]
        row_pekerjaan_ids = data["pages"][0]["row_pekerjaan_ids"]
        self.assertEqual(len(row_formulas), len(volume_rows))
        self.assertEqual(len(row_pekerjaan_ids), len(volume_rows))
        item_row_idx = next(idx for idx, row in enumerate(volume_rows) if len(row) >= 3 and row[0] == "1")
        self.assertEqual(row_formulas[item_row_idx], "=bp_1 * 2")
        self.assertEqual(row_pekerjaan_ids[item_row_idx], pekerjaan.id)

    @skipUnless(OPENPYXL_AVAILABLE, "openpyxl is required for XLSX export test")
    def test_xlsx_volume_column_uses_formula_reference_to_parameter_sheet(self):
        ProjectParameter.objects.create(
            project=self.project,
            name="bp_1",
            label="Panjang Dinding",
            value="7",
        )
        pekerjaan = self._create_one_pekerjaan()
        VolumeFormulaState.objects.create(
            project=self.project,
            pekerjaan=pekerjaan,
            raw="=bp_1 * 2",
            is_fx=True,
        )

        manager = ExportManager(self.project)
        response = manager.export_volume_pekerjaan("xlsx", parameters={"bp_1": 12})
        wb = load_workbook(BytesIO(response.content), data_only=False)

        self.assertIn("Parameters", wb.sheetnames)
        self.assertIn("Volume Pekerjaan", wb.sheetnames)

        ws_volume = wb["Volume Pekerjaan"]
        item_row = None
        for r in range(1, ws_volume.max_row + 1):
            if str(ws_volume.cell(row=r, column=1).value or "").strip() == "1":
                item_row = r
                break
        self.assertIsNotNone(item_row, "Row item nomor 1 tidak ditemukan pada sheet Volume Pekerjaan")

        volume_cell = ws_volume.cell(row=item_row, column=5)
        self.assertEqual(volume_cell.data_type, "f")
        self.assertTrue(str(volume_cell.value or "").startswith("=Parameters!$C$"))
        self.assertRegex(str(volume_cell.value or ""), r"\*\s*2")


class DeepCopyVolumeFormulaStateTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.owner = user_model.objects.create_user(
            username="deep_copy_formula_state_user",
            email="deep-copy-formula-state@example.com",
            password="Secret123!",
        )
        self.source = Project.objects.create(owner=self.owner, nama="Source Project Formula State")

    def test_copy_project_also_copies_and_remaps_volume_formula_state(self):
        ProjectParameter.objects.create(
            project=self.source,
            name="bp_10",
            label="Panjang",
            value="10",
        )
        ProjectComputedParameter.objects.create(
            project=self.source,
            name="cp_10",
            label="Luas",
            expression="=bp_10 * 2",
        )

        klas = Klasifikasi.objects.create(project=self.source, name="Klasifikasi")
        sub = SubKlasifikasi.objects.create(project=self.source, klasifikasi=klas, name="Sub")
        pekerjaan = Pekerjaan.objects.create(
            project=self.source,
            sub_klasifikasi=sub,
            source_type=Pekerjaan.SOURCE_CUSTOM,
            snapshot_kode="CUST-0001",
            snapshot_uraian="Pekerjaan Sumber",
            snapshot_satuan="m3",
        )
        VolumePekerjaan.objects.create(project=self.source, pekerjaan=pekerjaan, quantity="1")
        VolumeFormulaState.objects.create(
            project=self.source,
            pekerjaan=pekerjaan,
            raw="=bp_10 + cp_10",
            is_fx=True,
        )

        copied = DeepCopyService(self.source).copy(
            new_owner=self.owner,
            new_name="Copied Project Formula State",
            copy_jadwal=False,
        )

        copied_formula = VolumeFormulaState.objects.get(project=copied)
        copied_raw = copied_formula.raw or ""

        self.assertNotIn("bp_10", copied_raw.lower())
        self.assertNotIn("cp_10", copied_raw.lower())
        self.assertRegex(copied_raw.lower(), r"bp_[1-9][0-9]*")
        self.assertRegex(copied_raw.lower(), r"cp_[1-9][0-9]*")

        copied_param_names = set(
            ProjectParameter.objects.filter(project=copied).values_list("name", flat=True)
        )
        copied_computed_names = set(
            ProjectComputedParameter.objects.filter(project=copied).values_list("name", flat=True)
        )
        self.assertTrue(all(re.match(r"^bp_[1-9][0-9]*$", name) for name in copied_param_names))
        self.assertTrue(all(re.match(r"^cp_[1-9][0-9]*$", name) for name in copied_computed_names))
