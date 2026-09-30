"""WP Export — "dataset backend = nilai sel Excel" parity contract.

Owner decision (2026-06-20, all Option A): Excel is a *report* of the backend
SSOT, not a parallel calc engine. Money/quantity cells must therefore be real
Excel numbers equal to the backend canonical value (Decimal), at the agreed
precision — never a locale string and never a live formula.

This suite renders a small fixture through the real adapter+exporter and reads
the workbook back with openpyxl. One report per test class; add reports as each
slice lands (K2 -> K3 -> precision).
"""
import re
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.conf import settings

from openpyxl import load_workbook
from lxml import etree

from dashboard.models import Project
from detail_project.models import (
    Klasifikasi,
    SubKlasifikasi,
    Pekerjaan,
    HargaItemProject,
    DetailAHSPProject,
    DetailAHSPExpanded,
    VolumePekerjaan,
)
from detail_project.models import (
    ProjectParameter, ProjectComputedParameter, VolumeFormulaState, PekerjaanProgressWeekly,
)
from detail_project.exports.export_manager import ExportManager
from detail_project.exports.jadwal_pekerjaan_adapter import JadwalPekerjaanExportAdapter


class _RekapFixtureMixin:
    """Minimal project: 1 pekerjaan, koef 2 x harga 100 = E 200; +10% markup =>
    G 220; volume 3 => total 660."""

    def setUp(self):
        self.owner = get_user_model().objects.create_user(
            username="wp-export-parity-owner", password="not-used",
        )
        self.project = Project.objects.create(owner=self.owner, nama="Parity Project")
        klas = Klasifikasi.objects.create(project=self.project, name="Klas A", ordering_index=1)
        sub = SubKlasifikasi.objects.create(
            project=self.project, klasifikasi=klas, name="Sub A", ordering_index=1,
        )
        self.pekerjaan = Pekerjaan.objects.create(
            project=self.project, sub_klasifikasi=sub, source_type=Pekerjaan.SOURCE_CUSTOM,
            snapshot_kode="P-001", snapshot_uraian="Pekerjaan parity", snapshot_satuan="m2",
            ordering_index=1,
        )
        item = HargaItemProject.objects.create(
            project=self.project, kode_item="BHN-001", kategori="BHN",
            uraian="Bahan parity", satuan="kg", harga_satuan=Decimal("100.00"),
        )
        source = DetailAHSPProject.objects.create(
            project=self.project, pekerjaan=self.pekerjaan, harga_item=item,
            kategori="BHN", kode="BHN-001", uraian="Bahan parity", satuan="kg",
            koefisien=Decimal("2.000000"),
        )
        DetailAHSPExpanded.objects.create(
            project=self.project, pekerjaan=self.pekerjaan, source_detail=source,
            harga_item=item, kategori="BHN", kode="BHN-001", uraian="Bahan parity",
            satuan="kg", koefisien=Decimal("2.000000"), expansion_depth=0,
        )
        VolumePekerjaan.objects.create(
            project=self.project, pekerjaan=self.pekerjaan, quantity=Decimal("3.000"),
        )

    def _render_xlsx(self, response):
        return load_workbook(BytesIO(response.content))

    def _find_row(self, ws, col1_value):
        for row in ws.iter_rows():
            if row and row[0].value == col1_value:
                return row
        return None


class RekapRABExcelParityTests(_RekapFixtureMixin, TestCase):
    def test_money_cells_are_numbers_equal_to_backend(self):
        resp = ExportManager(self.project, self.owner).export_rekap_rab("xlsx")
        wb = self._render_xlsx(resp)
        ws = wb.worksheets[0]  # page 1 = RAB table
        row = self._find_row(ws, "Pekerjaan parity")
        self.assertIsNotNone(row, "pekerjaan row not found in RAB sheet")

        volume_cell, harga_cell, jumlah_cell = row[3], row[4], row[5]

        # Real numbers, not locale strings / formulas.
        for cell in (volume_cell, harga_cell, jumlah_cell):
            self.assertIsInstance(
                cell.value, (int, float),
                f"cell {cell.coordinate}={cell.value!r} is not a number",
            )
            self.assertFalse(str(cell.value).startswith("="), "cell holds a formula")

        # Values equal the backend canonical numbers (G=220, total=660, vol=3).
        self.assertAlmostEqual(float(volume_cell.value), 3.0, places=3)
        self.assertAlmostEqual(float(harga_cell.value), 220.0, places=2)
        self.assertAlmostEqual(float(jumlah_cell.value), 660.0, places=2)

        # Agreed precision (2 dp money, 3 dp volume).
        self.assertEqual(harga_cell.number_format, "#,##0.00")
        self.assertEqual(jumlah_cell.number_format, "#,##0.00")
        self.assertEqual(volume_cell.number_format, "#,##0.000")


class RekapRABWordParityTests(_RekapFixtureMixin, TestCase):
    """The text exporters (here: Word) must render the canonical Decimal as a
    properly formatted id-ID 2-dp string at their boundary — never the raw
    str(Decimal) '220.00' and never a live formula."""

    def test_word_cells_are_formatted_idid_strings(self):
        from docx import Document

        resp = ExportManager(self.project, self.owner).export_rekap_rab("word")
        doc = Document(BytesIO(resp.content))
        texts = [
            cell.text
            for table in doc.tables
            for row in table.rows
            for cell in row.cells
        ]
        # G=220, total=660 rendered id-ID with 2 decimals.
        self.assertIn("220,00", texts)
        self.assertIn("660,00", texts)
        # The raw Decimal repr must not leak through.
        self.assertNotIn("220.00", texts)
        self.assertNotIn("660.00", texts)


class RekapKebutuhanExcelParityTests(_RekapFixtureMixin, TestCase):
    """Same fixture: volume 3 x koef 2 => kebutuhan qty 6; harga 100; total 600.
    Qty is a 3-dp number, money is 2-dp, and the grand total reconciles."""

    def test_item_and_grand_total_are_numbers(self):
        resp = ExportManager(self.project, self.owner).export_rekap_kebutuhan("xlsx")
        ws = load_workbook(BytesIO(resp.content)).worksheets[0]

        # Item row: [No, Kode, Uraian, Satuan, Qty, Harga, Total]
        item_row = None
        for row in ws.iter_rows():
            if len(row) >= 7 and row[1].value == "BHN-001":
                item_row = row
                break
        self.assertIsNotNone(item_row, "kebutuhan item row not found")

        qty_cell, harga_cell, total_cell = item_row[4], item_row[5], item_row[6]
        for cell in (qty_cell, harga_cell, total_cell):
            self.assertIsInstance(
                cell.value, (int, float), f"{cell.coordinate}={cell.value!r} not numeric",
            )
        self.assertAlmostEqual(float(qty_cell.value), 6.0, places=3)
        self.assertAlmostEqual(float(harga_cell.value), 100.0, places=2)
        self.assertAlmostEqual(float(total_cell.value), 600.0, places=2)
        self.assertEqual(qty_cell.number_format, "#,##0.000")
        self.assertEqual(harga_cell.number_format, "#,##0.00")
        self.assertEqual(total_cell.number_format, "#,##0.00")

        # Grand Total footer reconciles to the sum of item totals, as a number.
        grand_cell = None
        for row in ws.iter_rows():
            if row and row[0].value == "Grand Total Harga":
                grand_cell = row[1]
                break
        if grand_cell is not None:
            self.assertIsInstance(grand_cell.value, (int, float))
            self.assertAlmostEqual(float(grand_cell.value), 600.0, places=2)
            self.assertEqual(grand_cell.number_format, "#,##0.00")


class RincianAHSPExcelParityTests(_RekapFixtureMixin, TestCase):
    """Component koef 2 x harga 100 => jumlah 200; E 200, +10% => F 20, G 220.
    Cells must be real numbers (no live =E*F / =SUM / cross-sheet formula)."""

    def test_no_formula_cells_and_values_match_backend(self):
        resp = ExportManager(self.project, self.owner).export_rincian_ahsp("xlsx")
        wb = load_workbook(BytesIO(resp.content))

        # Gate: not a single cell on the OFFICIAL sheets may be a live formula
        # (data_type 'f'). The "Kontrol Kalkulasi" audit sheet is excluded — it is
        # the control layer and intentionally carries formulas.
        for ws in wb.worksheets:
            if ws.title == "Kontrol Kalkulasi":
                continue
            for row in ws.iter_rows():
                for cell in row:
                    self.assertNotEqual(
                        cell.data_type, "f",
                        f"{ws.title}!{cell.coordinate} is a formula: {cell.value!r}",
                    )

        rincian = wb["Rincian"]
        # Detail row: [No, Uraian, Kode, Satuan, Koef, Harga, Jumlah]
        comp = None
        for row in rincian.iter_rows():
            if len(row) >= 7 and row[2].value == "BHN-001":
                comp = row
                break
        self.assertIsNotNone(comp, "component row not found")
        koef, harga, jumlah = comp[4], comp[5], comp[6]
        for cell in (koef, harga, jumlah):
            self.assertIsInstance(cell.value, (int, float))
        self.assertAlmostEqual(float(koef.value), 2.0, places=6)
        self.assertAlmostEqual(float(harga.value), 100.0, places=2)
        self.assertAlmostEqual(float(jumlah.value), 200.0, places=2)
        self.assertEqual(koef.number_format, "0.000000")
        self.assertEqual(harga.number_format, "#,##0.00")
        self.assertEqual(jumlah.number_format, "#,##0.00")

        # Rekap sheet row for the pekerjaan: [No, Kode, Uraian, E, F, G]
        rekap = wb["Rekap"]
        pek = None
        for row in rekap.iter_rows():
            if len(row) >= 6 and row[1].value == "P-001":
                pek = row
                break
        self.assertIsNotNone(pek, "pekerjaan recap row not found")
        e_cell, f_cell, g_cell = pek[3], pek[4], pek[5]
        for cell in (e_cell, f_cell, g_cell):
            self.assertIsInstance(cell.value, (int, float))
        self.assertAlmostEqual(float(e_cell.value), 200.0, places=2)
        self.assertAlmostEqual(float(f_cell.value), 20.0, places=2)
        self.assertAlmostEqual(float(g_cell.value), 220.0, places=2)

    def test_kontrol_kalkulasi_sheet_audits_via_formula(self):
        resp = ExportManager(self.project, self.owner).export_rincian_ahsp("xlsx")
        wb = load_workbook(BytesIO(resp.content))
        self.assertIn("Kontrol Kalkulasi", wb.sheetnames)
        kontrol = wb["Kontrol Kalkulasi"]

        # Row: [No, Kode, Uraian, Nilai Resmi, Nilai Kontrol, Selisih, Status]
        row = None
        for r in kontrol.iter_rows():
            if len(r) >= 7 and r[1].value == "P-001":
                row = r
                break
        self.assertIsNotNone(row, "kontrol row not found")
        resmi, kontrol_cell, selisih, status = row[3], row[4], row[5], row[6]

        # Nilai Resmi = backend G (numeric 220), not a formula.
        self.assertIsInstance(resmi.value, (int, float))
        self.assertAlmostEqual(float(resmi.value), 220.0, places=2)

        # The control columns are LIVE formulas (this is the audit layer).
        self.assertEqual(kontrol_cell.data_type, "f")
        self.assertIn("Rincian!", str(kontrol_cell.value))
        self.assertEqual(selisih.data_type, "f")
        self.assertEqual(status.data_type, "f")
        self.assertIn("OK", str(status.value))  # =IF(...,"OK","PERIKSA")

        # The control reconciles: the Rincian E + F cells the formula references sum
        # to the official G, so Selisih would be 0 and Status "OK".
        addrs = re.findall(r"Rincian!([A-Z]+\d+)", str(kontrol_cell.value))
        self.assertEqual(len(addrs), 2)
        rincian = wb["Rincian"]
        e_plus_f = sum(float(rincian[addr].value) for addr in addrs)
        self.assertAlmostEqual(e_plus_f, float(resmi.value), places=2)

    def test_word_rincian_uses_materialized_idid_strings(self):
        from docx import Document

        resp = ExportManager(self.project, self.owner).export_rincian_ahsp("word")
        doc = Document(BytesIO(resp.content))
        texts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    texts.append(cell.text)
        blob = "\n".join(texts)
        # G=220 rendered id-ID; the raw str(Decimal) must not leak.
        self.assertIn("220,00", blob)
        self.assertNotIn("220.00", blob)


class JadwalMonthlyValueOnlyTests(TestCase):
    """WP Export 2A: the Jadwal SSOT 'Data Master' sheet carries backend NUMBERS
    (no recompute formula), and the Monthly rincian sheet keeps only pure 1:1
    ='Data Master'!cell mirrors — no =SUM/arithmetic/multi-reference."""

    _MIRROR = re.compile(r"^='?Data Master'?!\$?[A-Za-z]+\$?\d+$")

    def setUp(self):
        self.owner = get_user_model().objects.create_user("wp-jadwal-2a-owner", password="x")
        self.project = Project.objects.create(
            owner=self.owner, nama="Jadwal 2A",
            tanggal_mulai=date(2026, 1, 1), tanggal_selesai=date(2026, 2, 11),
            week_start_day=0, week_end_day=6,
        )
        klas = Klasifikasi.objects.create(project=self.project, name="K", ordering_index=1)
        sub = SubKlasifikasi.objects.create(project=self.project, klasifikasi=klas, name="S", ordering_index=1)
        pkj = Pekerjaan.objects.create(
            project=self.project, sub_klasifikasi=sub, source_type=Pekerjaan.SOURCE_CUSTOM,
            snapshot_kode="P-001", snapshot_uraian="P", snapshot_satuan="m2", ordering_index=1,
        )
        item = HargaItemProject.objects.create(
            project=self.project, kode_item="BHN-1", kategori="BHN", uraian="B", satuan="kg",
            harga_satuan=Decimal("100.00"),
        )
        src = DetailAHSPProject.objects.create(
            project=self.project, pekerjaan=pkj, harga_item=item, kategori="BHN", kode="BHN-1",
            uraian="B", satuan="kg", koefisien=Decimal("2.000000"),
        )
        DetailAHSPExpanded.objects.create(
            project=self.project, pekerjaan=pkj, source_detail=src, harga_item=item, kategori="BHN",
            kode="BHN-1", uraian="B", satuan="kg", koefisien=Decimal("2.000000"), expansion_depth=0,
        )
        VolumePekerjaan.objects.create(project=self.project, pekerjaan=pkj, quantity=Decimal("10"))

        # A second pekerjaan with half the value and different progress locks the
        # weighted aggregation (the single-row/100%-weight case cannot catch mapping
        # mistakes between pekerjaan, planned, and actual rows).
        pkj_2 = Pekerjaan.objects.create(
            project=self.project, sub_klasifikasi=sub, source_type=Pekerjaan.SOURCE_CUSTOM,
            snapshot_kode="P-002", snapshot_uraian="P2", snapshot_satuan="m2", ordering_index=2,
        )
        src_2 = DetailAHSPProject.objects.create(
            project=self.project, pekerjaan=pkj_2, harga_item=item, kategori="BHN", kode="BHN-1",
            uraian="B", satuan="kg", koefisien=Decimal("1.000000"),
        )
        DetailAHSPExpanded.objects.create(
            project=self.project, pekerjaan=pkj_2, source_detail=src_2, harga_item=item, kategori="BHN",
            kode="BHN-1", uraian="B", satuan="kg", koefisien=Decimal("1.000000"), expansion_depth=0,
        )
        VolumePekerjaan.objects.create(project=self.project, pekerjaan=pkj_2, quantity=Decimal("10"))

        from datetime import timedelta
        for wk in range(1, 7):
            ws = date(2026, 1, 1) + timedelta(days=(wk - 1) * 7)
            PekerjaanProgressWeekly.objects.create(
                project=self.project, pekerjaan=pkj, week_number=wk,
                week_start_date=ws, week_end_date=ws + timedelta(days=6),
                planned_proportion=Decimal("15.00"),
                actual_proportion=Decimal("10.00"),
            )
            PekerjaanProgressWeekly.objects.create(
                project=self.project, pekerjaan=pkj_2, week_number=wk,
                week_start_date=ws, week_end_date=ws + timedelta(days=6),
                planned_proportion=Decimal("5.00"),
                actual_proportion=Decimal("20.00"),
            )

    @staticmethod
    def _value_for_label(ws, label, offset=2):
        for row in ws.iter_rows():
            for cell in row:
                if cell.value == label:
                    return ws.cell(row=cell.row, column=cell.column + offset)
        raise AssertionError(f"label {label!r} not found in {ws.title}")

    def test_data_master_numeric_and_monthly_rincian_mirror_only(self):
        resp = ExportManager(self.project, self.owner).export_jadwal_professional(
            "xlsx", report_type="monthly", months=[1],
        )
        wb = load_workbook(BytesIO(resp.content))
        self.assertIn("Data Master", wb.sheetnames)

        # SSOT gate: not one cell on Data Master is a live formula.
        for cell in (c for row in wb["Data Master"].iter_rows() for c in row):
            self.assertNotEqual(
                cell.data_type, "f",
                f"Data Master!{cell.coordinate} is a formula: {cell.value!r}",
            )

        # Monthly rincian gate: every formula is a pure 1:1 ='Data Master'!cell mirror.
        rincian = next(
            (wb[s] for s in wb.sheetnames if "Rincian" in s or s.startswith("M")), None,
        )
        self.assertIsNotNone(rincian, f"monthly rincian sheet not found in {wb.sheetnames}")
        formula_cells = 0
        for cell in (c for row in rincian.iter_rows() for c in row):
            if cell.data_type == "f":
                formula_cells += 1
                self.assertRegex(
                    str(cell.value), self._MIRROR,
                    f"{rincian.title}!{cell.coordinate} is not a 1:1 mirror: {cell.value!r}",
                )
        self.assertGreater(formula_cells, 0, "expected mirror formulas on the rincian sheet")

    def test_padded_final_period_preserves_weighted_planned_and_actual_values(self):
        """Period 2 carries data in W5-W6 and pads W7-W8 with zero. Aggregates
        must preserve both weighted planned and actual values through the padding."""
        resp = ExportManager(self.project, self.owner).export_jadwal_professional(
            "xlsx", report_type="monthly", months=[2],
        )
        wb = load_workbook(BytesIO(resp.content), data_only=False)
        rincian = next(
            (wb[s] for s in wb.sheetnames if "Rincian" in s or s.startswith("M")), None,
        )
        self.assertIsNotNone(rincian, f"monthly rincian sheet not found in {wb.sheetnames}")

        # Values are in a 2:1 ratio, so weights are 2/3 and 1/3.
        # Planned per week=(2/3*15% + 1/3*5%)=7/60; actual=(2/3*10% + 1/3*20%)=2/15.
        self.assertAlmostEqual(self._value_for_label(rincian, "Rencana Bulan Ini").value, 7 / 30, places=9)
        self.assertAlmostEqual(self._value_for_label(rincian, "Realisasi Bulan Ini").value, 4 / 15, places=9)

        # Cumulative remains a pure SSOT mirror through the padded W7-W8 columns.
        cumulative = self._value_for_label(rincian, "Kumulatif s.d Ini")
        self.assertEqual(cumulative.data_type, "f")
        self.assertRegex(str(cumulative.value), self._MIRROR)

        # Through W6: weighted actual=80%, planned=70%, therefore deviation=+10%.
        self.assertAlmostEqual(self._value_for_label(rincian, "Deviasi").value, 0.10, places=9)

        # Per-pekerjaan monthly planned values remain unweighted in rows (30% and
        # 10%); the TOTAL row applies the 2:1 project weights and yields 7/30.
        col_i_values = [
            cell.value for cell in (rincian.cell(row=r, column=9) for r in range(1, rincian.max_row + 1))
            if isinstance(cell.value, (int, float))
        ]
        self.assertTrue(any(abs(value - 0.30) < 1e-9 for value in col_i_values))
        self.assertTrue(any(abs(value - 0.10) < 1e-9 for value in col_i_values))
        self.assertTrue(any(abs(value - (7 / 30)) < 1e-9 for value in col_i_values))

    def test_weekly_rincian_is_values_with_mirror_only_formulas(self):
        """WP Export 2B: Weekly rincian carries Python-weighted VALUES; the only
        formulas left are pure 1:1 ='Data Master'!cell mirrors. Per-pekerjaan
        Progress Minggu Ini = bobot×planned (2/3×15%=10% and 1/3×5%=1/60); the
        TOTAL row is the weighted project value 7/60."""
        resp = ExportManager(self.project, self.owner).export_jadwal_professional(
            "xlsx", report_type="weekly", weeks=[1],
        )
        wb = load_workbook(BytesIO(resp.content), data_only=False)

        # Data Master is still numeric-only (shared SSOT).
        for cell in (c for row in wb["Data Master"].iter_rows() for c in row):
            self.assertNotEqual(cell.data_type, "f", f"Data Master!{cell.coordinate}={cell.value!r}")

        rincian = next(
            (wb[s] for s in wb.sheetnames if "Rincian" in s or "Minggu" in s), None,
        )
        self.assertIsNotNone(rincian, f"weekly rincian sheet not found in {wb.sheetnames}")
        formula_cells = 0
        for cell in (c for row in rincian.iter_rows() for c in row):
            if cell.data_type == "f":
                formula_cells += 1
                self.assertRegex(
                    str(cell.value), self._MIRROR,
                    f"{rincian.title}!{cell.coordinate} is not a 1:1 mirror: {cell.value!r}",
                )
        self.assertGreater(formula_cells, 0, "expected mirror formulas on the weekly rincian sheet")

        # Col I (Progress Minggu Ini) = bobot × planned per pekerjaan + weighted TOTAL.
        col_i = [
            cell.value for cell in (rincian.cell(row=r, column=9) for r in range(1, rincian.max_row + 1))
            if isinstance(cell.value, (int, float))
        ]
        self.assertTrue(any(abs(v - 0.10) < 1e-9 for v in col_i), col_i)      # P-001: 2/3×15%
        self.assertTrue(any(abs(v - (1 / 60)) < 1e-9 for v in col_i), col_i)  # P-002: 1/3×5%
        self.assertTrue(any(abs(v - (7 / 60)) < 1e-9 for v in col_i), col_i)  # TOTAL (weighted)

        # Ringkasan rencana tetap nilai planned (K-12 menambah blok realisasi
        # terpisah, tidak mengubah angka rencana).
        self.assertAlmostEqual(
            self._value_for_label(rincian, "Progress Kumulatif s.d. Minggu Lalu", 5).value,
            0,
            places=9,
        )
        self.assertAlmostEqual(
            self._value_for_label(rincian, "Progress Minggu Ini", 5).value,
            7 / 60,
            places=9,
        )
        self.assertAlmostEqual(
            self._value_for_label(rincian, "Progress Kumulatif s.d. Minggu Ini", 5).value,
            7 / 60,
            places=9,
        )

        total_row = next(
            r for r in range(1, rincian.max_row + 1) if rincian.cell(row=r, column=1).value == "TOTAL"
        )
        self.assertAlmostEqual(rincian.cell(total_row, 6).value, 3300.00, places=9)
        self.assertAlmostEqual(rincian.cell(total_row, 7).value, 1.00, places=9)
        self.assertAlmostEqual(rincian.cell(total_row, 8).value, 0.00, places=9)
        self.assertAlmostEqual(rincian.cell(total_row, 9).value, 7 / 60, places=9)
        self.assertAlmostEqual(rincian.cell(total_row, 10).value, 7 / 60, places=9)

        # K-12 (owner 2026-09-30): realisasi tampil untuk SEMUA proyek, juga
        # proyek tanpa masa tambahan. Actual/minggu = 2/3*10% + 1/3*20% = 2/15.
        self.assertTrue(any(
            cell.value == "REALISASI MINGGU KE-1"
            for row in rincian.iter_rows() for cell in row
        ))
        self.assertAlmostEqual(
            self._value_for_label(rincian, "Realisasi Kumulatif s.d. Minggu Lalu", 5).value, 0, places=9,
        )
        self.assertAlmostEqual(
            self._value_for_label(rincian, "Realisasi Minggu Ini", 5).value, 2 / 15, places=9,
        )
        self.assertAlmostEqual(
            self._value_for_label(rincian, "Realisasi Kumulatif s.d. Minggu Ini", 5).value, 2 / 15, places=9,
        )
        actual_total_row = next(
            r for r in range(1, rincian.max_row + 1)
            if rincian.cell(row=r, column=1).value == "TOTAL REALISASI"
        )
        self.assertAlmostEqual(rincian.cell(actual_total_row, 9).value, 2 / 15, places=9)

    @staticmethod
    def _assert_no_formulas(ws):
        for cell in (c for row in ws.iter_rows() for c in row):
            if cell.data_type == "f":
                raise AssertionError(f"{ws.title}!{cell.coordinate}={cell.value!r}")

    def test_2c_professional_kurva_uses_backend_values_and_chart_remains_valid(self):
        resp = ExportManager(self.project, self.owner).export_jadwal_professional(
            "xlsx", report_type="rekap",
        )
        wb = load_workbook(BytesIO(resp.content), data_only=False)
        kurva = wb["Kurva S"]
        gantt = wb["Input Progress-Gantt"]

        self._assert_no_formulas(kurva)
        self._assert_no_formulas(gantt)
        self.assertEqual(len(kurva._charts), 1)
        self.assertEqual(len(kurva._charts[0].series), 2)

        # Six weeks: planned=(2/3*15% + 1/3*5%)*6=70%;
        # actual=(2/3*10% + 1/3*20%)*6=80%.
        planned_row = self._value_for_label(kurva, "Kumulatif Rencana", 0).row
        actual_row = self._value_for_label(kurva, "Kumulatif Realisasi", 0).row
        final_week_col = 7 + 6  # G is W0, H-M are W1-W6.
        self.assertAlmostEqual(kurva.cell(planned_row, final_week_col).value, 0.70, places=9)
        self.assertAlmostEqual(kurva.cell(actual_row, final_week_col).value, 0.80, places=9)

        first_week_col = 8
        detail_planned = []
        detail_actual = []
        for row in range(4, planned_row):
            planned_value = kurva.cell(row, first_week_col).value
            actual_value = kurva.cell(row + 1, first_week_col).value
            if (
                isinstance(kurva.cell(row, 6).value, (int, float))
                and kurva.cell(row, 6).value > 0
                and isinstance(planned_value, (int, float))
                and isinstance(actual_value, (int, float))
            ):
                detail_planned.append(planned_value)
                detail_actual.append(actual_value)
        self.assertAlmostEqual(sum(detail_planned), 7 / 60, places=9)
        self.assertAlmostEqual(sum(detail_actual), 2 / 15, places=9)
        self.assertAlmostEqual(kurva.cell(planned_row - 2, first_week_col).value, 7 / 60, places=9)
        self.assertAlmostEqual(kurva.cell(planned_row - 1, first_week_col).value, 2 / 15, places=9)

        cover = wb["Cover"]
        formulas = [
            c for row in cover.iter_rows() for c in row if c.data_type == "f"
        ]
        self.assertEqual(len(formulas), 2)
        for cell in formulas:
            self.assertRegex(str(cell.value), r"^='Kurva S'![A-Z]+\d+$")
        self.assertAlmostEqual(self._value_for_label(cover, "Deviasi").value, 0.10, places=9)

    def test_2c_monthly_shared_kurva_uses_backend_values_and_chart(self):
        resp = ExportManager(self.project, self.owner).export_jadwal_professional(
            "xlsx", report_type="monthly", months=[1],
        )
        wb = load_workbook(BytesIO(resp.content), data_only=False)
        kurva = wb["Kurva S M1"]
        self._assert_no_formulas(kurva)
        self.assertEqual(len(kurva._charts), 1)
        self.assertEqual(len(kurva._charts[0].series), 2)

        planned_row = self._value_for_label(kurva, "Kumulatif Rencana", 0).row
        actual_row = self._value_for_label(kurva, "Kumulatif Realisasi", 0).row
        final_week_col = 7 + 4  # G is W0, H-K are W1-W4.
        self.assertAlmostEqual(kurva.cell(planned_row, final_week_col).value, 7 / 15, places=9)
        self.assertAlmostEqual(kurva.cell(actual_row, final_week_col).value, 8 / 15, places=9)

    def test_2c_zero_project_total_produces_zero_weights_without_formulas(self):
        HargaItemProject.objects.filter(project=self.project).update(harga_satuan=Decimal("0"))
        resp = ExportManager(self.project, self.owner).export_jadwal_professional(
            "xlsx", report_type="rekap",
        )
        wb = load_workbook(BytesIO(resp.content), data_only=False)
        kurva = wb["Kurva S"]
        self._assert_no_formulas(kurva)

        total_row = next(
            r for r in range(1, kurva.max_row + 1)
            if kurva.cell(r, 2).value == "TOTAL"
        )
        self.assertEqual(kurva.cell(total_row, 6).value, 0)
        self.assertEqual(kurva.cell(total_row, 7).value, 0)
        for label in ("Progress Mingguan Rencana", "Progress Mingguan Realisasi",
                      "Kumulatif Rencana", "Kumulatif Realisasi"):
            row = self._value_for_label(kurva, label, 0).row
            for col in range(7, 7 + 6 + 1):
                self.assertEqual(kurva.cell(row, col).value, 0)

    def test_2d_backend_parity_and_stable_identity_across_jadwal_reports(self):
        # Duplicate uraian proves the professional path no longer joins rows by name.
        Pekerjaan.objects.filter(project=self.project, snapshot_kode="P-002").update(
            snapshot_uraian="P"
        )
        backend = JadwalPekerjaanExportAdapter(self.project).get_rekap_report_data()
        self.assertEqual(backend["meta"]["total_pekerjaan"], 2)
        expected_planned = backend["kurva_s_data"][-1]["planned"] / 100
        expected_actual = backend["kurva_s_data"][-1]["actual"] / 100

        professional = load_workbook(BytesIO(
            ExportManager(self.project, self.owner).export_jadwal_professional(
                "xlsx", report_type="rekap",
            ).content
        ), data_only=False)
        kurva = professional["Kurva S"]
        input_progress = professional["Input Progress-Gantt"]
        self.assertEqual(
            sum(1 for row in input_progress.iter_rows() if isinstance(row[0].value, int)),
            2,
        )
        self.assertEqual(
            sorted(c.value for c in kurva["A"] if c.value in {"P-001", "P-002"}),
            ["P-001", "P-002"],
        )
        prof_planned_row = self._value_for_label(kurva, "Kumulatif Rencana", 0).row
        prof_actual_row = self._value_for_label(kurva, "Kumulatif Realisasi", 0).row
        self.assertAlmostEqual(kurva.cell(prof_planned_row, 13).value, expected_planned, places=9)
        self.assertAlmostEqual(kurva.cell(prof_actual_row, 13).value, expected_actual, places=9)

        monthly = load_workbook(BytesIO(
            ExportManager(self.project, self.owner).export_jadwal_professional(
                "xlsx", report_type="monthly", months=[2],
            ).content
        ), data_only=False)
        monthly_master = monthly["Data Master"]
        monthly_kurva = monthly["Kurva S M2"]
        self.assertEqual(
            sum(1 for row in monthly_master.iter_rows() if isinstance(row[0].value, int)),
            2,
        )
        master_planned_row = self._value_for_label(monthly_master, "Kumulatif Rencana", 0).row
        master_actual_row = self._value_for_label(monthly_master, "Kumulatif Realisasi", 0).row
        self.assertAlmostEqual(monthly_master.cell(master_planned_row, 13).value, expected_planned, places=9)
        self.assertAlmostEqual(monthly_master.cell(master_actual_row, 13).value, expected_actual, places=9)
        month_planned_row = self._value_for_label(monthly_kurva, "Kumulatif Rencana", 0).row
        month_actual_row = self._value_for_label(monthly_kurva, "Kumulatif Realisasi", 0).row
        self.assertAlmostEqual(monthly_kurva.cell(month_planned_row, 13).value, expected_planned, places=9)
        self.assertAlmostEqual(monthly_kurva.cell(month_actual_row, 13).value, expected_actual, places=9)

        weekly = load_workbook(BytesIO(
            ExportManager(self.project, self.owner).export_jadwal_professional(
                "xlsx", report_type="weekly", weeks=[6],
            ).content
        ), data_only=False)
        weekly_master = weekly["Data Master"]
        weekly_rincian = weekly["Rincian Progress W6"]
        weekly_planned_row = self._value_for_label(weekly_master, "Kumulatif Rencana", 0).row
        weekly_actual_row = self._value_for_label(weekly_master, "Kumulatif Realisasi", 0).row
        self.assertAlmostEqual(weekly_master.cell(weekly_planned_row, 13).value, expected_planned, places=9)
        self.assertAlmostEqual(weekly_master.cell(weekly_actual_row, 13).value, expected_actual, places=9)
        self.assertAlmostEqual(
            self._value_for_label(
                weekly_rincian, "Progress Kumulatif s.d. Minggu Ini", 5
            ).value,
            expected_planned,
            places=9,
        )


class JadwalDailyDocxExportTests(TestCase):
    """Regression guard for the template-based daily DOCX export.

    The documentation page can contain SmartArt. Copying that page repeatedly
    must keep WordprocessingML drawing IDs and diagram relationships unique, or
    Microsoft Word refuses to open the generated document.
    """

    NS = {
        "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
        "wp14": "http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing",
        "dgm": "http://schemas.openxmlformats.org/drawingml/2006/diagram",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    }

    def setUp(self):
        self.owner = get_user_model().objects.create_user("wp-jadwal-daily-owner", password="x")
        self.project = Project.objects.create(
            owner=self.owner,
            nama="Daily DOCX",
            sumber_dana="APBD",
            lokasi_project="Lokasi",
            nama_client="Owner Harian",
            anggaran_owner=Decimal("1000000.00"),
            nama_kontraktor="Kontraktor Harian",
            nama_konsultan_pengawas="Pengawas Harian",
            tanggal_mulai=date(2026, 1, 1),
            tanggal_selesai=date(2026, 1, 28),
            week_start_day=0,
            week_end_day=6,
        )
        klas = Klasifikasi.objects.create(project=self.project, name="K", ordering_index=1)
        sub = SubKlasifikasi.objects.create(project=self.project, klasifikasi=klas, name="S", ordering_index=1)
        self.pekerjaan = Pekerjaan.objects.create(
            project=self.project,
            sub_klasifikasi=sub,
            source_type=Pekerjaan.SOURCE_CUSTOM,
            snapshot_kode="P-001",
            snapshot_uraian="Pekerjaan harian",
            snapshot_satuan="m2",
            ordering_index=1,
        )
        item = HargaItemProject.objects.create(
            project=self.project,
            kode_item="BHN-1",
            kategori="BHN",
            uraian="B",
            satuan="kg",
            harga_satuan=Decimal("100.00"),
        )
        src = DetailAHSPProject.objects.create(
            project=self.project,
            pekerjaan=self.pekerjaan,
            harga_item=item,
            kategori="BHN",
            kode="BHN-1",
            uraian="B",
            satuan="kg",
            koefisien=Decimal("1.000000"),
        )
        DetailAHSPExpanded.objects.create(
            project=self.project,
            pekerjaan=self.pekerjaan,
            source_detail=src,
            harga_item=item,
            kategori="BHN",
            kode="BHN-1",
            uraian="B",
            satuan="kg",
            koefisien=Decimal("1.000000"),
            expansion_depth=0,
        )
        VolumePekerjaan.objects.create(project=self.project, pekerjaan=self.pekerjaan, quantity=Decimal("10"))

        from datetime import timedelta
        for wk in range(1, 5):
            start = date(2026, 1, 1) + timedelta(days=(wk - 1) * 7)
            PekerjaanProgressWeekly.objects.create(
                project=self.project,
                pekerjaan=self.pekerjaan,
                week_number=wk,
                week_start_date=start,
                week_end_date=start + timedelta(days=6),
                planned_proportion=Decimal("10.00"),
                actual_proportion=Decimal("5.00"),
            )

    def _daily_docx_content(self, daily_mode="month", period=1):
        response = ExportManager(self.project, self.owner).export_jadwal_professional(
            "word",
            report_type="daily",
            daily_mode=daily_mode,
            period=period,
        )
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.assertTrue(response.content.startswith(b"PK"))
        return response.content, response["Content-Disposition"]

    def _assert_unique_docx_drawing_ids(self, content):
        with ZipFile(BytesIO(content)) as package:
            self.assertIsNone(package.testzip())
            names = set(package.namelist())
            self.assertIn("word/document.xml", names)
            self.assertIn("word/_rels/document.xml.rels", names)
            self.assertIn("[Content_Types].xml", names)

            doc = etree.fromstring(package.read("word/document.xml"))
            docpr_ids = [node.get("id") for node in doc.xpath(".//wp:docPr", namespaces=self.NS)]
            anchor_ids = [value for value in doc.xpath(".//@wp14:anchorId", namespaces=self.NS)]
            edit_ids = [value for value in doc.xpath(".//@wp14:editId", namespaces=self.NS)]

            for values, label in (
                (docpr_ids, "wp:docPr id"),
                (anchor_ids, "wp14:anchorId"),
                (edit_ids, "wp14:editId"),
            ):
                self.assertEqual(len(values), len(set(values)), f"duplicate {label}")

            rel_ids = []
            for node in doc.xpath(".//dgm:relIds", namespaces=self.NS):
                for attr in ("dm", "lo", "qs", "cs"):
                    value = node.get(f"{{{self.NS['r']}}}{attr}")
                    if value:
                        rel_ids.append(value)
            self.assertEqual(len(rel_ids), len(set(rel_ids)), "duplicate SmartArt relationship IDs")

            diagram_parts = [name for name in names if name.startswith("word/diagrams/")]
            if rel_ids:
                self.assertGreater(len(diagram_parts), 0)
                content_types = package.read("[Content_Types].xml").decode("utf-8")
                self.assertIn("/word/diagrams/", content_types)

    def test_daily_docx_template_exists_for_deploy(self):
        template = Path(settings.BASE_DIR) / "detail_project" / "export_templates" / "laporan_harian_template.docx"
        self.assertTrue(template.exists(), f"missing daily DOCX template: {template}")
        with ZipFile(template) as package:
            self.assertIsNone(package.testzip())
            self.assertIn("word/document.xml", package.namelist())

    def test_daily_docx_month_mode_has_unique_smartart_parts_and_filename_range(self):
        content, disposition = self._daily_docx_content("month", 1)
        self.assertIn('filename="Laporan Harian 01-01 - 25-01.docx"', disposition)
        self._assert_unique_docx_drawing_ids(content)

    def test_daily_docx_rejects_non_word_format_at_manager_boundary(self):
        with self.assertRaises(ValueError):
            ExportManager(self.project, self.owner).export_jadwal_professional(
                "xlsx",
                report_type="daily",
                daily_mode="day",
                period=1,
            )

    def _daily_docx(self, day):
        from docx import Document

        response = ExportManager(self.project, self.owner).export_jadwal_professional(
            "word", report_type="daily", daily_mode="day", days=[day],
        )
        return Document(BytesIO(response.content))

    @staticmethod
    def _table_texts(table):
        return [cell.text for row in table.rows for cell in row.cells]

    def test_daily_progress_is_cumulative_up_to_previous_week(self):
        # Fixture: rencana 10%/minggu, realisasi 5%/minggu, 1 pekerjaan (bobot 1).
        # 15 Jan = Minggu 3 -> kumulatif W1..W2 = 20% / 10% / -10%, bukan 10/5/-5
        # (bug lama: hanya minggu sebelumnya saja).
        texts = self._table_texts(self._daily_docx(15).tables[0])
        self.assertIn("s.d. Minggu 2", texts)
        self.assertIn("20.00%", texts)
        self.assertIn("10.00%", texts)
        self.assertIn("-10.00%", texts)

    def test_daily_current_week_shows_planned_target_only(self):
        # 15 Jan = Minggu 3: target kumulatif W1..W3 = 30% (+10% minggu ini);
        # realisasi & deviasi minggu berjalan belum ada -> '-'.
        table = self._daily_docx(15).tables[0]
        current_col = [row.cells[4].text for row in table.rows]
        self.assertEqual(current_col, ["s.d. Minggu 3", "30.00% (+10.00%)", "-", "-"])

    def test_daily_progress_first_week_shows_zero_not_dash(self):
        table = self._daily_docx(1).tables[0]
        previous_col = [row.cells[3].text for row in table.rows]
        self.assertEqual(previous_col, ["Awal Proyek", "0.00%", "0.00%", "0.00%"])
        self.assertEqual(table.rows[1].cells[4].text, "10.00% (+10.00%)")

    def test_daily_identity_table_drops_fields_shown_elsewhere(self):
        texts = " ".join(self._table_texts(self._daily_docx(15).tables[0]))
        for removed in ("No. Kontrak", "Kontraktor :", "Konsultan :", "Tgl Laporan"):
            self.assertNotIn(removed, texts)
        for kept in ("Proyek :", "Lokasi :", "Cuaca :", "Pemilik/Penanggung Jawab Project :"):
            self.assertIn(kept, texts)

    def test_daily_signatures_contractor_and_supervisor_only_with_signing_space(self):
        from docx.shared import Cm

        doc = self._daily_docx(15)
        signatures = doc.tables[2]
        # R-37: baris 0 = instansi (kosong di fixture), baris 2 = nama; tanpa sebutan peran.
        names = [cell.text for cell in signatures.rows[2].cells]
        self.assertEqual(names, ["Kontraktor Harian", "Pengawas Harian"])
        flat = " ".join(self._table_texts(signatures))
        for stale in ("Pemilik", "Owner Harian", "Kontraktor Pelaksana", "Konsultan Pengawas"):
            self.assertNotIn(stale, flat)
        self.assertGreaterEqual(signatures.rows[1].height, Cm(2.1))  # 2.2cm, dibulatkan ke twips

    def test_daily_rejects_dates_outside_project_with_user_message(self):
        from detail_project.exports.errors import ExportValidationError, export_error_response

        # Proyek 01-28 Jan 2026: hari ke-40 di luar masa proyek.
        with self.assertRaises(ExportValidationError) as ctx:
            ExportManager(self.project, self.owner).export_jadwal_professional(
                "word", report_type="daily", daily_mode="day", days=[40],
            )
        response = export_error_response(ctx.exception)
        self.assertEqual(response.status_code, 400)
        self.assertIn("di luar masa proyek", response.content.decode("utf-8"))

    def test_daily_rejects_project_without_start_date(self):
        from detail_project.exports.errors import ExportValidationError

        Project.objects.filter(pk=self.project.pk).update(tanggal_mulai=None)
        self.project.refresh_from_db()
        with self.assertRaises(ExportValidationError):
            ExportManager(self.project, self.owner).export_jadwal_professional(
                "word", report_type="daily", daily_mode="day", days=[1],
            )


class VolumeExportParityTests(TestCase):
    """Volume value = canonical stored quantity (numeric); base param = backend
    value (numeric); computed param = expression text + Nilai '-'; no live formula.
    Verifies the PDF/Word boundary (materialize) for the Volume pages structure."""

    def setUp(self):
        self.owner = get_user_model().objects.create_user(
            username="wp-volume-parity-owner", password="not-used",
        )
        self.project = Project.objects.create(owner=self.owner, nama="Volume Parity")
        klas = Klasifikasi.objects.create(project=self.project, name="Klas A", ordering_index=1)
        sub = SubKlasifikasi.objects.create(
            project=self.project, klasifikasi=klas, name="Sub A", ordering_index=1,
        )
        self.pekerjaan = Pekerjaan.objects.create(
            project=self.project, sub_klasifikasi=sub, source_type=Pekerjaan.SOURCE_CUSTOM,
            snapshot_kode="P-001", snapshot_uraian="Pekerjaan volume", snapshot_satuan="m3",
            ordering_index=1,
        )
        VolumePekerjaan.objects.create(
            project=self.project, pekerjaan=self.pekerjaan, quantity=Decimal("125.500"),
        )
        ProjectParameter.objects.create(
            project=self.project, name="bp_1", label="Panjang", value=Decimal("7"),
        )
        ProjectComputedParameter.objects.create(
            project=self.project, name="cp_1", label="Luas", expression="bp_1 * 2",
        )
        VolumeFormulaState.objects.create(
            project=self.project, pekerjaan=self.pekerjaan, raw="=cp_1 + bp_1", is_fx=True,
        )

    def test_word_volume_and_params_are_materialized_numbers(self):
        from docx import Document

        resp = ExportManager(self.project, self.owner).export_volume_pekerjaan("word")
        doc = Document(BytesIO(resp.content))
        blob = "\n".join(
            [p.text for p in doc.paragraphs]
            + [c.text for t in doc.tables for r in t.rows for c in r.cells]
        )
        # Canonical volume 125.5 -> id-ID 3 dp; raw str(Decimal) must not leak.
        self.assertIn("125,500", blob)
        self.assertNotIn("125.500", blob)
        # Base param backend value present (7 -> "7,00"); computed Nilai is "-".
        self.assertIn("7,00", blob)


class HargaItemsExcelParityTests(TestCase):
    """Three USED price items: one priced, one NULL (not filled), one explicit 0.00.
    The Excel price cell must be a real 2-dp number for priced/zero, and a distinct
    '-' marker for NULL (D-HI-01: NULL must never collapse into 0.00)."""

    def setUp(self):
        self.owner = get_user_model().objects.create_user(
            username="wp-harga-parity-owner", password="not-used",
        )
        self.project = Project.objects.create(owner=self.owner, nama="Harga Parity")
        klas = Klasifikasi.objects.create(project=self.project, name="Klas A", ordering_index=1)
        sub = SubKlasifikasi.objects.create(
            project=self.project, klasifikasi=klas, name="Sub A", ordering_index=1,
        )
        pekerjaan = Pekerjaan.objects.create(
            project=self.project, sub_klasifikasi=sub, source_type=Pekerjaan.SOURCE_CUSTOM,
            snapshot_kode="P-001", snapshot_uraian="Pekerjaan harga", snapshot_satuan="m2",
            ordering_index=1,
        )
        # kode, harga_satuan
        specs = [
            ("BHN-PRICED", Decimal("1500.00")),
            ("BHN-NULLED", None),
            ("BHN-ZEROED", Decimal("0.00")),
        ]
        for kode, harga in specs:
            item = HargaItemProject.objects.create(
                project=self.project, kode_item=kode, kategori="BHN",
                uraian=f"Bahan {kode}", satuan="kg", harga_satuan=harga,
            )
            source = DetailAHSPProject.objects.create(
                project=self.project, pekerjaan=pekerjaan, harga_item=item,
                kategori="BHN", kode=kode, uraian=f"Bahan {kode}", satuan="kg",
                koefisien=Decimal("1.000000"),
            )
            DetailAHSPExpanded.objects.create(
                project=self.project, pekerjaan=pekerjaan, source_detail=source,
                harga_item=item, kategori="BHN", kode=kode, uraian=f"Bahan {kode}",
                satuan="kg", koefisien=Decimal("1.000000"), expansion_depth=0,
            )

    def _harga_cell(self, ws, kode):
        for row in ws.iter_rows():
            if len(row) >= 5 and row[1].value == kode:
                return row[4]
        return None

    def test_price_cells_numeric_and_null_stays_distinct(self):
        resp = ExportManager(self.project, self.owner).export_harga_items("xlsx")
        ws = load_workbook(BytesIO(resp.content)).worksheets[0]

        priced = self._harga_cell(ws, "BHN-PRICED")
        nulled = self._harga_cell(ws, "BHN-NULLED")
        zeroed = self._harga_cell(ws, "BHN-ZEROED")
        for cell, kode in ((priced, "PRICED"), (nulled, "NULLED"), (zeroed, "ZEROED")):
            self.assertIsNotNone(cell, f"row {kode} not found")

        # Priced + explicit zero are real numbers at 2 dp.
        self.assertIsInstance(priced.value, (int, float))
        self.assertAlmostEqual(float(priced.value), 1500.0, places=2)
        self.assertEqual(priced.number_format, "#,##0.00")

        self.assertIsInstance(zeroed.value, (int, float))
        self.assertAlmostEqual(float(zeroed.value), 0.0, places=2)

        # NULL stays a distinct text marker — never a numeric 0.
        self.assertEqual(nulled.value, "-")
        self.assertNotEqual(nulled.value, zeroed.value)
