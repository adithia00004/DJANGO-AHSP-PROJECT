"""Fixtures and rendered-output checks for export after the contract end.

Type 1 extends the final week only. Type 2 extends that week and adds W7.
Both scenarios use the same canonical Monday-Sunday week boundary as project 217.
"""

import base64
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from docx import Document

from dashboard.models import Project
from detail_project.exports.errors import ExportValidationError
from detail_project.exports.export_manager import ExportManager
from detail_project.exports.jadwal_pekerjaan_adapter import JadwalPekerjaanExportAdapter
from detail_project.exports.pdf_exporter import PDFExporter
from detail_project.models import (
    DetailAHSPExpanded,
    DetailAHSPProject,
    HargaItemProject,
    Klasifikasi,
    Pekerjaan,
    PekerjaanProgressWeekly,
    SubKlasifikasi,
    VolumePekerjaan,
)
from detail_project.tests_jadwal_monthly_report import pdf_page_texts
from detail_project.timeline_utils import (
    contract_boundary_week,
    is_extension_day,
    is_extension_week,
    project_report_period_counts,
    work_period_end,
)


class ExtensionExportFixtureMixin:
    @classmethod
    def setUpTestData(cls):
        cls.owner = get_user_model().objects.create_user(
            username="export-extension-owner", password="StrongPass123!"
        )
        base = {
            "owner": cls.owner,
            "sumber_dana": "APBD",
            "lokasi_project": "Makassar",
            "nama_client": "Pemilik",
            "anggaran_owner": Decimal("1000000.00"),
            "week_start_day": 0,
            "week_end_day": 6,
        }
        cls.same_week = Project.objects.create(
            **base,
            nama="Tambahan dalam minggu batas",
            tanggal_mulai=date(2026, 9, 14),
            tanggal_selesai=date(2026, 9, 16),  # Rabu
            tanggal_akhir_tambahan=date(2026, 9, 19),  # Sabtu
        )
        cls.next_week = Project.objects.create(
            **base,
            nama="Tambahan sampai W7",
            tanggal_mulai=date(2026, 8, 10),
            tanggal_selesai=date(2026, 9, 19),  # Sabtu di W6
            tanggal_akhir_tambahan=date(2026, 9, 27),  # Minggu di W7
        )
        cls.without_extension = Project.objects.create(
            **base,
            nama="Tanpa tambahan",
            tanggal_mulai=date(2026, 8, 10),
            tanggal_selesai=date(2026, 9, 19),
        )


class ExtensionExportFixtureTests(ExtensionExportFixtureMixin, TestCase):
    def test_same_week_extension_has_days_but_no_new_week(self):
        self.assertEqual(work_period_end(self.same_week), date(2026, 9, 19))
        self.assertEqual(contract_boundary_week(self.same_week), 1)
        self.assertEqual(project_report_period_counts(self.same_week), (1, 1))
        self.assertFalse(is_extension_week(self.same_week, 1))
        self.assertFalse(is_extension_day(self.same_week, date(2026, 9, 16)))
        self.assertTrue(is_extension_day(self.same_week, date(2026, 9, 17)))

    def test_new_week_extension_keeps_boundary_in_w6(self):
        self.assertEqual(work_period_end(self.next_week), date(2026, 9, 27))
        self.assertEqual(contract_boundary_week(self.next_week), 6)
        self.assertEqual(project_report_period_counts(self.next_week), (7, 2))
        self.assertFalse(is_extension_week(self.next_week, 6))
        self.assertTrue(is_extension_week(self.next_week, 7))
        self.assertTrue(is_extension_day(self.next_week, date(2026, 9, 20)))

    def test_project_without_extension_has_no_markers(self):
        self.assertEqual(work_period_end(self.without_extension), date(2026, 9, 19))
        self.assertEqual(project_report_period_counts(self.without_extension), (6, 2))
        self.assertFalse(is_extension_week(self.without_extension, 7))
        self.assertFalse(is_extension_day(self.without_extension, date(2026, 9, 20)))


class ExtensionExportAdapterTests(ExtensionExportFixtureMixin, TestCase):
    def test_same_week_extension_marks_boundary_without_an_extra_week(self):
        data = JadwalPekerjaanExportAdapter(self.same_week).get_rekap_report_data()
        self.assertEqual(data["contract_end"], date(2026, 9, 16))
        self.assertEqual(data["additional_end"], date(2026, 9, 19))
        self.assertEqual(data["boundary_week"], 1)
        self.assertEqual(len(data["weekly_columns"]), 1)
        self.assertTrue(data["weekly_columns"][0]["is_boundary_week"])
        self.assertFalse(data["weekly_columns"][0]["is_extension_week"])

    def test_w6_boundary_and_w7_extension_reach_all_report_payloads(self):
        adapter = JadwalPekerjaanExportAdapter(self.next_week)
        for data in (
            adapter.get_rekap_report_data(),
            adapter.get_monthly_comparison_data(2),
            adapter.get_weekly_comparison_data(7),
        ):
            self.assertEqual(data["contract_end"], date(2026, 9, 19))
            self.assertEqual(data["additional_end"], date(2026, 9, 27))
            self.assertEqual(data["boundary_week"], 6)
        columns = adapter.get_rekap_report_data()["weekly_columns"]
        self.assertEqual(len(columns), 7)
        self.assertTrue(columns[5]["is_boundary_week"])
        self.assertFalse(columns[5]["is_extension_week"])
        self.assertFalse(columns[6]["is_boundary_week"])
        self.assertTrue(columns[6]["is_extension_week"])

    def test_project_without_extension_has_no_column_markers(self):
        data = JadwalPekerjaanExportAdapter(self.without_extension).get_rekap_report_data()
        self.assertIsNone(data["additional_end"])
        self.assertIsNone(data["boundary_week"])
        self.assertEqual(len(data["weekly_columns"]), 6)
        self.assertTrue(all(not col["is_boundary_week"] for col in data["weekly_columns"]))
        self.assertTrue(all(not col["is_extension_week"] for col in data["weekly_columns"]))


class ExtensionDailyWordTests(ExtensionExportFixtureMixin, TestCase):
    def _document_text(self, project, week):
        response = ExportManager(project, self.owner).export_jadwal_professional(
            "word", report_type="daily", daily_mode="week", period=week,
        )
        self.assertTrue(response.content.startswith(b"PK"))
        document = Document(BytesIO(response.content))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)

    def test_w7_daily_word_renders_extension_days(self):
        text = self._document_text(self.next_week, 7)
        self.assertIn("Senin, 21 September 2026 | Minggu 7 | Penambahan Waktu Kerja", text)
        self.assertNotIn("terlambat", text.lower())

    def test_boundary_week_marks_only_the_day_after_contract_end(self):
        text = self._document_text(self.next_week, 6)
        self.assertIn("Sabtu, 19 September 2026 | Minggu 6", text)
        self.assertNotIn("Sabtu, 19 September 2026 | Minggu 6 | Penambahan", text)
        self.assertIn("Minggu, 20 September 2026 | Minggu 6 | Penambahan Waktu Kerja", text)

    def test_same_week_extension_marks_days_without_a_new_week(self):
        text = self._document_text(self.same_week, 1)
        self.assertIn("Rabu, 16 September 2026 | Minggu 1", text)
        self.assertIn("Kamis, 17 September 2026 | Minggu 1 | Penambahan Waktu Kerja", text)

    def test_day_after_work_period_remains_rejected(self):
        with self.assertRaises(ExportValidationError):
            ExportManager(self.next_week, self.owner).export_jadwal_professional(
                "word", report_type="daily", daily_mode="day", days=[50],
            )


class ExtensionSummaryFixtureMixin(ExtensionExportFixtureMixin):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        klasifikasi = Klasifikasi.objects.create(
            project=cls.next_week, name="Konstruksi", ordering_index=1,
        )
        sub = SubKlasifikasi.objects.create(
            project=cls.next_week, klasifikasi=klasifikasi, name="Pekerjaan", ordering_index=1,
        )
        for number, (name, unit_price, actual_w6, actual_w7) in enumerate((
            ("Pekerjaan A", "100", "100", "0"),
            ("Pekerjaan B", "300", "50", "50"),
        ), start=1):
            pekerjaan = Pekerjaan.objects.create(
                project=cls.next_week, sub_klasifikasi=sub,
                source_type=Pekerjaan.SOURCE_CUSTOM,
                snapshot_kode=f"P-{number:03d}", snapshot_uraian=name,
                snapshot_satuan="m2", ordering_index=number,
            )
            volume = Decimal("10")
            VolumePekerjaan.objects.create(
                project=cls.next_week, pekerjaan=pekerjaan, quantity=volume,
            )
            item = HargaItemProject.objects.create(
                project=cls.next_week, kode_item=f"BHN-{number}", kategori="BHN",
                uraian=name, satuan="m2", harga_satuan=Decimal(unit_price),
            )
            detail = DetailAHSPProject.objects.create(
                project=cls.next_week, pekerjaan=pekerjaan, harga_item=item,
                kategori="BHN", kode=item.kode_item, uraian=name,
                satuan="m2", koefisien=Decimal("1"),
            )
            DetailAHSPExpanded.objects.create(
                project=cls.next_week, pekerjaan=pekerjaan, source_detail=detail,
                harga_item=item, kategori="BHN", kode=item.kode_item,
                uraian=name, satuan="m2", koefisien=Decimal("1"), expansion_depth=0,
            )
            for week, start, end, planned, actual in (
                (6, date(2026, 9, 14), date(2026, 9, 20), "100", actual_w6),
                (7, date(2026, 9, 21), date(2026, 9, 27), "0", actual_w7),
            ):
                PekerjaanProgressWeekly.objects.create(
                    project=cls.next_week, pekerjaan=pekerjaan,
                    week_number=week, week_start_date=start, week_end_date=end,
                    planned_proportion=Decimal(planned), actual_proportion=Decimal(actual),
                )


class ExtensionSummaryDataTests(ExtensionSummaryFixtureMixin, TestCase):
    def test_summary_matches_weighted_weekly_progress_at_contract_boundary(self):
        adapter = JadwalPekerjaanExportAdapter(self.next_week)
        summary = adapter.get_contract_end_summary()
        weekly = adapter.get_weekly_comparison_data(6)

        self.assertEqual(summary["boundary_week"], 6)
        self.assertEqual((summary["boundary_start"], summary["boundary_end"]), (
            date(2026, 9, 14), date(2026, 9, 20),
        ))
        self.assertAlmostEqual(float(summary["planned"]), 100.0)
        self.assertAlmostEqual(float(summary["actual"]), 62.5)
        self.assertAlmostEqual(float(summary["actual"]), weekly["current_data"]["cumulative_actual"])
        self.assertEqual(len(summary["unfinished"]), 1)
        unfinished = summary["unfinished"][0]
        self.assertEqual(unfinished["description"], "Pekerjaan B")
        self.assertAlmostEqual(float(unfinished["weight"]), 75.0)
        self.assertAlmostEqual(float(unfinished["actual"]), 50.0)
        self.assertAlmostEqual(float(unfinished["remaining"]), 50.0)
        self.assertAlmostEqual(float(unfinished["remaining_weight"]), 37.5)


class ExtensionSummaryWordTests(ExtensionSummaryFixtureMixin, TestCase):
    def _word(self, project, **selection):
        response = ExportManager(project, self.owner).export_jadwal_professional(
            "word", report_type="daily", **selection,
        )
        return Document(BytesIO(response.content))

    def test_summary_separates_contract_day_from_extension_day(self):
        document = self._word(self.next_week, daily_mode="day", days=[41, 42])
        paragraphs = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertLess(paragraphs.index("Sabtu, 19 September 2026"),
                        paragraphs.index("RANGKUMAN PROGRESS AKHIR WAKTU KERJA"))
        self.assertLess(paragraphs.index("RANGKUMAN PROGRESS AKHIR WAKTU KERJA"),
                        paragraphs.index("Minggu, 20 September 2026"))
        self.assertIn("minggu batas 14/09/2026–20/09/2026", paragraphs)
        table_text = " ".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        self.assertIn("62,50%", table_text)
        self.assertIn("37,50%", table_text)
        summary_table = next(
            table for table in document.tables
            if table.cell(0, 0).text == "No" and table.cell(0, 1).text == "Uraian"
        )
        self.assertEqual(len(summary_table.rows), 2)
        self.assertEqual(summary_table.cell(1, 1).text, "Pekerjaan B")

    def test_extension_only_selection_starts_with_summary(self):
        document = self._word(self.next_week, daily_mode="week", period=7)
        paragraphs = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertLess(paragraphs.index("RANGKUMAN PROGRESS AKHIR WAKTU KERJA"),
                        paragraphs.index("Senin, 21 September 2026"))

    def test_same_week_extension_places_summary_before_thursday(self):
        document = self._word(self.same_week, daily_mode="week", period=1)
        paragraphs = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertLess(paragraphs.index("Rabu, 16 September 2026"),
                        paragraphs.index("RANGKUMAN PROGRESS AKHIR WAKTU KERJA"))
        self.assertLess(paragraphs.index("RANGKUMAN PROGRESS AKHIR WAKTU KERJA"),
                        paragraphs.index("Kamis, 17 September 2026"))

    def test_before_boundary_or_without_extension_has_no_summary(self):
        before = self._word(self.next_week, daily_mode="week", period=5)
        normal = self._word(self.without_extension, daily_mode="week", period=6)
        for document in (before, normal):
            self.assertNotIn(
                "RANGKUMAN PROGRESS AKHIR WAKTU KERJA",
                "\n".join(paragraph.text for paragraph in document.paragraphs),
            )


class ExtensionSummaryPdfTests(ExtensionSummaryFixtureMixin, TestCase):
    def _pages(self, report_type, **selection):
        response = ExportManager(self.next_week, self.owner).export_jadwal_professional(
            "pdf", report_type=report_type, **selection,
        )
        self.assertTrue(response.content.startswith(b"%PDF"))
        return pdf_page_texts(response.content)

    def test_weekly_boundary_then_summary_then_extension_week(self):
        pages = self._pages("weekly", weeks=[6, 7])
        boundary = next(i for i, page in enumerate(pages) if "PROGRESS PELAKSANAAN PEKERJAAN MINGGU KE-6" in page)
        summary = next(i for i, page in enumerate(pages) if "RANGKUMAN PROGRESS AKHIR WAKTU KERJA" in page)
        signature = next(i for i, page in enumerate(pages) if i > summary and "LEMBAR PENGESAHAN" in page)
        extension = next(i for i, page in enumerate(pages) if "PROGRESS PELAKSANAAN PEKERJAAN MINGGU KE-7" in page)
        self.assertLess(boundary, summary)
        self.assertLess(summary, signature)
        self.assertLess(signature, extension)
        self.assertNotIn("PROGRESS PELAKSANAAN PEKERJAAN", pages[summary])
        self.assertIn("62,50%", pages[summary])
        self.assertIn("Pekerjaan B", pages[summary])
        self.assertNotIn("Pekerjaan A", pages[summary])

    def test_extension_only_starts_with_summary_before_weekly_cover(self):
        pages = self._pages("weekly", weeks=[7])
        summary = next(i for i, page in enumerate(pages) if "RANGKUMAN PROGRESS AKHIR WAKTU KERJA" in page)
        week_seven = next(i for i, page in enumerate(pages) if "PROGRESS PELAKSANAAN PEKERJAAN MINGGU KE-7" in page)
        self.assertLess(summary, week_seven)

    def test_monthly_boundary_summary_follows_month_report(self):
        pages = self._pages("monthly", months=[1, 2])
        month_two = next(i for i, page in enumerate(pages) if "PROGRESS PELAKSANAAN PEKERJAAN BULAN KE-2" in page)
        summary = next(i for i, page in enumerate(pages) if "RANGKUMAN PROGRESS AKHIR WAKTU KERJA" in page)
        signature = next(i for i, page in enumerate(pages) if i > summary and "LEMBAR PENGESAHAN" in page)
        self.assertLess(month_two, summary)
        self.assertLess(summary, signature)

    def test_same_week_extension_has_summary_but_no_extension_week(self):
        response = ExportManager(self.same_week, self.owner).export_jadwal_professional(
            "pdf", report_type="weekly", weeks=[1],
        )
        pages = pdf_page_texts(response.content)
        self.assertEqual(sum("RANGKUMAN PROGRESS AKHIR WAKTU KERJA" in page for page in pages), 1)
        self.assertFalse(any("MINGGU KE-2" in page for page in pages))

    def test_single_period_pdf_paths_place_summary_once(self):
        for report_type, period in (("weekly", 6), ("monthly", 2)):
            with self.subTest(report_type=report_type):
                pages = self._pages(report_type, period=period)
                self.assertEqual(
                    sum("RANGKUMAN PROGRESS AKHIR WAKTU KERJA" in page for page in pages),
                    1,
                )
                summary = next(i for i, page in enumerate(pages) if "RANGKUMAN PROGRESS AKHIR WAKTU KERJA" in page)
                signature = next(i for i, page in enumerate(pages) if i > summary and "LEMBAR PENGESAHAN" in page)
                self.assertLess(summary, signature)

    def test_before_boundary_and_without_extension_have_no_summary(self):
        before = self._pages("weekly", weeks=[5])
        response = ExportManager(self.without_extension, self.owner).export_jadwal_professional(
            "pdf", report_type="weekly", weeks=[6],
        )
        for pages in (before, pdf_page_texts(response.content)):
            self.assertFalse(any("RANGKUMAN PROGRESS AKHIR WAKTU KERJA" in page for page in pages))


class ExtensionPdfMarkerTests(ExtensionSummaryFixtureMixin, TestCase):
    def test_rekap_boundary_line_is_after_w6_not_w7(self):
        manager = ExportManager(self.next_week, self.owner)
        page = JadwalPekerjaanExportAdapter(self.next_week).get_rekap_report_data()['planned_pages'][0]
        config = manager._create_config_simple('Jadwal', page_orientation='landscape', page_size='A3')
        table = PDFExporter(config)._build_table(page)
        boundary_lines = [
            command for command in table._linecmds
            if command[0] == 'LINEAFTER' and command[3] == 2.5
        ]
        self.assertEqual(len(boundary_lines), 1)
        self.assertEqual(boundary_lines[0][1], (8, 0))  # 3 static + W6 at index 5

    def test_rekap_planned_and_actual_grids_label_additional_week(self):
        response = ExportManager(self.next_week, self.owner).export_jadwal_professional(
            "pdf", report_type="rekap",
        )
        pages = pdf_page_texts(response.content)
        planned = next(page for page in pages if "GRID VIEW - RENCANA" in page)
        actual = next(page for page in pages if "GRID VIEW - REALISASI" in page)
        for page in (planned, actual):
            self.assertIn("W6", page)
            self.assertIn("W7", page)
            self.assertIn("Penambahan", page)

    def test_gantt_uses_server_boundary_when_frontend_sends_columns(self):
        pekerjaan = list(Pekerjaan.objects.filter(project=self.next_week).order_by('id'))
        gantt_data = {
            'rows': [
                {'id': item.id, 'name': item.snapshot_uraian, 'type': 'pekerjaan',
                 'volume': 10, 'satuan': 'm2', 'level': 3}
                for item in pekerjaan
            ],
            'time_columns': [
                {'week': week, 'label': f'W{week}', 'range': ''}
                for week in range(1, 8)
            ],
            'planned': {str(item.id): {6: 100} for item in pekerjaan},
            'actual': {str(item.id): {6: 50, 7: 50} for item in pekerjaan},
        }
        response = ExportManager(self.next_week, self.owner).export_jadwal_professional(
            "pdf", report_type="rekap", gantt_data=gantt_data,
        )
        pages = pdf_page_texts(response.content)
        gantt = next(page for page in pages if "BAGIAN 4: GANTT CHART" in page)
        self.assertIn("Penambahan", gantt)

    def test_monthly_kurva_tables_mark_the_additional_column(self):
        response = ExportManager(self.next_week, self.owner).export_jadwal_professional(
            "pdf", report_type="monthly", months=[2],
        )
        pages = pdf_page_texts(response.content)
        kurva_pages = [
            page for page in pages
            if "RINGKASAN PROGRESS KURVA S" in page or "GRAFIK KURVA S" in page
        ]
        self.assertTrue(kurva_pages)
        self.assertTrue(all("Penambahan" in page for page in kurva_pages))

    def test_rekap_kurva_table_marks_w7(self):
        response = ExportManager(self.next_week, self.owner).export_jadwal_professional(
            "pdf", report_type="rekap",
        )
        pages = pdf_page_texts(response.content)
        kurva_pages = [page for page in pages if "KURVA S" in page and "W7" in page]
        self.assertTrue(kurva_pages)
        self.assertTrue(any("Penambahan" in page for page in kurva_pages))

    def test_browser_kurva_attachment_does_not_duplicate_server_kurva(self):
        manager = ExportManager(self.next_week, self.owner)
        base = pdf_page_texts(manager.export_jadwal_professional("pdf", report_type="rekap").content)
        tiny_png = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/lXcAAAAASUVORK5CYII="
        )
        attached = pdf_page_texts(manager.export_jadwal_professional(
            "pdf", report_type="rekap",
            attachments=[{'title': 'Kurva S dari browser', 'bytes': tiny_png}],
        ).content)
        self.assertEqual(len(attached), len(base))
        self.assertEqual(sum("KURVA S" in page for page in attached),
                         sum("KURVA S" in page for page in base))

    def test_weekly_extension_subtitle_is_absent_from_boundary_week(self):
        response = ExportManager(self.next_week, self.owner).export_jadwal_professional(
            "pdf", report_type="weekly", weeks=[6, 7],
        )
        pages = pdf_page_texts(response.content)
        week_six = next(page for page in pages if "PROGRESS PELAKSANAAN PEKERJAAN MINGGU KE-6" in page)
        week_seven = next(page for page in pages if "PROGRESS PELAKSANAAN PEKERJAAN MINGGU KE-7" in page)
        self.assertNotIn("Penambahan Waktu Kerja", week_six)
        self.assertIn("Penambahan Waktu Kerja", week_seven)

    def test_later_month_has_subtitle_and_front_summary(self):
        self.next_week.tanggal_akhir_tambahan = date(2026, 10, 11)  # W9, month 3
        self.next_week.save(update_fields=['tanggal_akhir_tambahan'])
        response = ExportManager(self.next_week, self.owner).export_jadwal_professional(
            "pdf", report_type="monthly", months=[3],
        )
        pages = pdf_page_texts(response.content)
        summary = next(i for i, page in enumerate(pages) if "RANGKUMAN PROGRESS AKHIR WAKTU KERJA" in page)
        month_three = next(i for i, page in enumerate(pages) if "PROGRESS PELAKSANAAN PEKERJAAN BULAN KE-3" in page)
        self.assertLess(summary, month_three)
        self.assertIn("Penambahan Waktu Kerja", pages[month_three])
